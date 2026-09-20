"""ta_calc_ledger -- 계산 장부 + 엑셀 모델 + 본문 대조 테스트 (Task 10).

손계산 가능한 작은 재무 픽스처로 값을 고정한다. 네트워크 없음.

실행: python tests/test_ta_calc_ledger.py
"""
import io
import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'scripts'))

import openpyxl                                          # noqa: E402
import ta_calc_ledger as cl                              # noqa: E402

if (getattr(sys.stdout, 'encoding', '') or '').lower().replace('-', '') != 'utf8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

_passed = 0
_failed = []


def eq(got, want, label):
    global _passed
    if got == want:
        _passed += 1
    else:
        _failed.append((label, want, got))


def near(got, want, label, tol=1e-6):
    ok = got is not None and want is not None and abs(got - want) <= tol
    eq(ok, True, f'{label} (got={got}, want={want})')


FS = {
    'meta': {'stock_name': '테스트'},
    'kis': {'market_cap_억': 1200},
    'forward': {'year_actual': '2025/12'},
    'financials': {
        '2024': {'revenue': 1000, 'op_income': 100, 'net_income': 80, 'eps': 800.0,
                 'total_equity': 500, 'total_assets': 1000, 'total_debt': 500,
                 'current_assets': 400, 'current_liabilities': 200,
                 'retained_earnings': 300, 'net_debt': 100},
        '2025': {'revenue': 1100, 'op_income': 123.1, 'net_income': 99, 'eps': 1077.2,
                 'total_equity': 600, 'total_assets': 1200, 'total_debt': 600,
                 'current_assets': 450, 'current_liabilities': 300,
                 'retained_earnings': 360, 'net_debt': 100},
        # 반기 실적이 섞인 미완결 연도 -- 최근 회계연도로 쓰면 안 된다
        '2026': {'revenue': 500, 'op_income': 10},
    },
}
MARKET = {'latest': {'close': 10000}, 'var': {'var95_1d': -0.034, 'var95_20d': -0.12, 'n': 250}}


def items_by(ledger, key, period_sub=None):
    return [i for i in ledger['items'] if i['key'] == key
            and (period_sub is None or period_sub in i['period'])]


def one(ledger, key, period_sub=None):
    xs = items_by(ledger, key, period_sub)
    return xs[0] if xs else None


# ---------- 비율 / 성장 / 레버리지 ----------
L = cl.build_ledger(FS, MARKET, None)
eq(L['fy'], '2025', 'forward.year_actual 로 최근 회계연도 결정 (2026 미완결 제외)')
r25 = one(L, 'roe', '2025')
near(r25['value'], 99 / ((500 + 600) / 2), 'roe 2025 = 순이익 / 평균 자본 (지배주주 분리 없음 -> 연결)')
eq('연결 기준' in r25.get('note', ''), True, '지배주주 분리 없으면 note 에 연결 기준 명시')
eq(r25['definition'], 'ROE = 연결 순이익 / 평균 자본총계((기초+기말)/2)', 'roe definition 기록')
r24 = one(L, 'roe', '2024')
near(r24['value'], 80 / 500, 'roe 2024: 기초(2023) 없으면 기말 기준')
eq('기말 기준' in r24['note'] and '기말' in r24['definition'], True, '기말 fallback 을 note/definition 에 명시')
near(one(L, 'roa', '2025')['value'], 99 / ((1000 + 1200) / 2), 'roa 2025 = 순이익 / 평균 자산')
eq(bool(one(L, 'opm', '2025')['definition']), True, '모든 비율 항목에 definition')
lr = one(L, 'dol')
eq(lr['status'], 'ok', 'dol ok')

# ---------- fallback: _fnguide / DART ----------
FN = {'financial': {'annual': {'net_income_controlling': {'2025/12': '90', '2024/12': '70'},
                               'finance_cost': {'2025/12': '10.0'}}},
      'consensus_estimates': {'_meta': {'actual': {'controlling_equity':
                                                   {'2023/12': 400, '2024/12': 480, '2025/12': 520}}}}}
DART = {'2025': {'list': [
    {'sj_div': 'BS', 'account_id': 'ifrs-full_RetainedEarnings', 'account_detail': '-', 'thstrm_amount': '36000000000'},
    {'sj_div': 'SCE', 'account_id': 'ifrs-full_RetainedEarnings', 'account_detail': '-', 'thstrm_amount': '1'},
    {'sj_div': 'CIS', 'account_id': 'ifrs-full_FinanceCosts', 'account_detail': '-', 'thstrm_amount': '2000000000'},
]}}
fs_nore = {**FS, 'financials': {**FS['financials'],
                                '2025': {k: v for k, v in FS['financials']['2025'].items() if k != 'retained_earnings'}}}
LF = cl.build_ledger(fs_nore, MARKET, None, fnguide=FN, dart=DART)
rf = one(LF, 'roe', '2025')
near(rf['value'], 90 / ((480 + 520) / 2), 'ROE = 지배주주 순이익 / 평균 지배주주지분 (fnguide fallback)')
eq(rf['definition'], 'ROE = 지배주주 순이익 / 평균 지배주주지분((기초+기말)/2)', '지배주주 정의')
eq(rf.get('note'), None, '지배주주 분리 있으면 연결 note 없음')
eq(rf['inputs']['net_income_ctrl_2025']['source'],
   '_fnguide.json:financial.annual.net_income_controlling.2025/12', '입력 출처 file:key')
eq(one(LF, 'altman_z')['status'], 'ok', 'altman ok (이익잉여금 fallback 후)')
eq(one(LF, 'altman_z')['inputs']['retained_earnings_2025']['source'].startswith('data_dart_financials.json:2025'),
   True, '이익잉여금 DART fallback (BS 만, SCE 무시)')
near(one(LF, 'altman_z')['inputs']['retained_earnings_2025']['value'], 360.0, 'DART 원 -> 억원')
ic = one(LF, 'interest_coverage', '2025')
near(ic['value'], 123.1 / 20, '이자보상배율: DART 금융원가 대용 (fnguide 보다 DART 우선)')
eq('금융원가' in ic['definition'] and '대용' in ic['note'], True, '금융원가 대용을 definition/note 에 명시')

# ---------- 정밀도 우선순위: DART(원) -> FnGuide(억 소수) -> financial_summary(정수 억) ----------
FS_R = {'forward': {'year_actual': '2025/12'}, 'kis': {},
        'financials': {'2024': {'net_income': -16, 'total_assets': 5061},
                       '2025': {'net_income': -143, 'total_assets': 5401}}}
FN_R = {'financial': {'annual': {'net_income': {'2024/12': '-16.38', '2025/12': '-143.3'}}},
        'consensus_estimates': {'_meta': {'actual': {'total_assets': {'2024/12': 5061.44, '2025/12': 5400.7}}}}}
DART_R = {'2025': {'list': [{'sj_div': 'BS', 'account_id': 'ifrs-full_Assets', 'account_detail': '-',
                             'thstrm_amount': '540070062623'}]}}
LR = cl.build_ledger(FS_R, None, None, fnguide=FN_R, dart=DART_R)
ra = one(LR, 'roa', '2025')
src = {k: v['source'].split(':')[0] for k, v in ra['inputs'].items()}
eq(src, {'net_income_2025': '_fnguide.json', 'total_assets_2025': 'data_dart_financials.json',
         'total_assets_2024': '_fnguide.json'}, 'DART > FnGuide > financial_summary 순으로 입력 선택')
near(ra['value'], -143.3 / ((5061.44 + 5400.70062623) / 2), 'ROA 는 비반올림 입력으로 계산', 1e-12)
eq(round(ra['value'] * 100, 2), -2.74, '에프에스티형 ROA 가 공표치 -2.74% 와 소수 2자리 일치')
near(one(cl.build_ledger(FS_R, None, None), 'roa', '2025')['value'], -143 / 5231, '다른 파일 없으면 financial_summary', 1e-12)
eq(ra['input_unit'].startswith('억원'), True, '항목마다 입력 단위 명시')


# ---------- 조용한 fallback 경고 + 연결/별도 기준 ----------
eq(any(w.startswith('total_assets(2025): DART 계정 없음') for w in LR['warnings']), False,
   'DART 에서 찾은 필드는 경고 없음')
eq('net_income(2025): DART 계정 없음 -> _fnguide.json:financial.annual.net_income.2025/12 사용' in LR['warnings'],
   True, 'DART -> FnGuide fallback 경고')
eq(any(w.startswith('total_assets(2025): DART 파일 없음 -> financial_summary.json')
       for w in cl.build_ledger(FS_R, None, None)['warnings']), True, 'DART 파일 자체가 없을 때도 경고')

eq(cl.source_basis({'dart': DART_R}, 'data_dart_financials.json:2025.list[BS:ifrs-full_Assets]'), 'unknown',
   'DART: 지배/비지배 계정도 fs_div 도 없으면 unknown (별도로 단정하지 않음)')
D_CFS = {'2025': {'list': [{'sj_div': 'BS', 'account_id': 'ifrs-full_NoncontrollingInterests', 'thstrm_amount': '1'}]}}
eq(cl.source_basis({'dart': D_CFS}, 'data_dart_financials.json:2025.list'), 'consolidated', 'DART: 비지배 계정 -> 연결')
D_OFS = {'2025': {'list': [{'sj_div': 'BS', 'account_id': 'x', 'fs_div': 'OFS', 'thstrm_amount': '1'}]}}
eq(cl.source_basis({'dart': D_OFS}, 'data_dart_financials.json:2025.list'), 'separate', 'DART: fs_div=OFS -> 별도')
eq(cl.source_basis({'fnguide': FN_R}, '_fnguide.json:x'), 'unknown', 'FnGuide: 지배주주 계정 없으면 unknown')
eq(cl.source_basis({'fnguide': FN}, '_fnguide.json:x'), 'consolidated', 'FnGuide: 지배주주 계정 -> 연결')
eq(cl.source_basis({'fs': FS}, 'financial_summary.json:kis.market_cap_억'), None, '시가총액은 기준 판정 대상 아님')

# 혼재: DART(unknown) + FnGuide(unknown) 로 계산한 ROA -> 경고
eq(ra.get('warnings'), [cl.MIXED_BASIS_WARNING], '두 파일 기준 unknown -> 연결/별도 혼재 경고')
eq(ra['basis'], 'unknown', 'basis=unknown')
# 같은 기준(연결) 두 파일 -> 경고 없음
DART_C = {'2025': {'list': DART_R['2025']['list'] + D_CFS['2025']['list']}}
FN_C = {**FN_R, 'financial': {'annual': {**FN_R['financial']['annual'], 'net_income_controlling': {}}}}  # 지배주주 계정 존재 = 연결 표지
rc = one(cl.build_ledger(FS_R, None, None, fnguide=FN_C, dart=DART_C), 'roa', '2025')
eq((rc['basis'], rc.get('warnings')), ('consolidated', None), '모두 연결이면 basis=consolidated, 경고 없음')
# 연결 DART + 별도 표시 financial_summary -> 경고
FS_SEP = {**FS_R, '_sources': {'balance_sheet': '별도 재무상태표'}}
rs = one(cl.build_ledger(FS_SEP, None, None, dart=DART_C), 'roa', '2025')
eq((rs['basis'], rs.get('warnings')), ('unknown', [cl.MIXED_BASIS_WARNING]), '연결+별도 섞이면 경고')
# 한 파일뿐이면 unknown 이어도 혼재 아님
r1 = one(cl.build_ledger(FS_R, None, None), 'roa', '2025')
eq((r1['basis'], r1.get('warnings')), ('unknown', None), '단일 출처는 혼재 경고 없음')

sk = one(L, 'interest_coverage', '2025')
eq('_fnguide.json' in sk['reason'] and 'data_dart_financials.json' in sk['reason'], True,
   'skip 사유에 확인한 파일 목록')
near(one(L, 'opm', '2025')['value'], 123.1 / 1100, 'opm 2025')
near(one(L, 'debt_ratio', '2025')['value'], 1.0, 'debt_ratio 2025 = 부채/자본')
near(one(L, 'current_ratio', '2025')['value'], 1.5, 'current_ratio 2025')
near(one(L, 'dol')['value'], 2.31, 'dol')
near(one(L, 'dfl')['value'], 1.5, 'dfl')
near(one(L, 'dcl')['value'], 2.31 * 1.5, 'dcl = dol x dfl')
near(one(L, 'sales_growth')['value'], 0.1, 'sales_growth')
eq(one(L, 'interest_coverage', '2025')['status'], 'skipped', '이자비용 없으면 skipped (0 으로 채우지 않음)')
eq(one(L, 'rim_value_per_share')['status'], 'skipped', 'assumptions 없으면 RIM skipped')
near(one(L, 'var95_1d')['value'], -0.034, 'VaR 는 market_data 에서 그대로')
eq('market_data.json' in str(one(L, 'var95_1d')['inputs']), True, 'VaR 출처 표기')

# 레버리지 가드
v, why = cl.leverage_ratio(100, 110, 1000, 1005)
eq((v, why), (None, '분모 변화 미미'), 'dol 가드: 분모 변화 <1%')
v, why = cl.leverage_ratio(-10, 20, 1000, 1100)
eq((v, why), (None, '부호 전환으로 해석 불가'), 'dol 가드: 전년 적자')
v, why = cl.leverage_ratio(23, -5, 2374, 2803)
eq((v, '흑자->적자' in why), (None, True), 'dol 확장 가드: 당년 적자 전환은 사유에 명시')
v, why = cl.leverage_ratio(123.1, 150, 100, 123.1)
eq(why, '', '정상 입력은 사유 없음')

# ---------- Altman Z ----------
z = one(L, 'altman_z')
want_z = (1.2 * (150 / 1200) + 1.4 * (360 / 1200) + 3.3 * (123.1 / 1200)
          + 0.6 * (1200 / 600) + 1.0 * (1100 / 1200))
near(z['value'], want_z, 'Altman Z 5요소 (B17 식)')
eq(z['zone'], 'safe', 'Z 3.03 -> safe')
eq(z.get('note'), cl.ALTMAN_NOTE, '순차입 기업은 시점/모형 note 만')
eq("시총은 현재 시점, 재무는 FY 기말" in z['note'] and "Z'' 가 적합" in z['note'], True, 'Altman note 문구')
eq([cl.altman_zone(x) for x in (3.0, 2.99, 1.81, 1.8)], ['safe', 'grey', 'grey', 'distress'], 'Z 구간 경계')
fs_cash = {**FS, 'financials': {**FS['financials'], '2025': {**FS['financials']['2025'], 'net_debt': -50}}}
z2 = one(cl.build_ledger(fs_cash, MARKET, None), 'altman_z')
eq(z2['status'], 'ok', '순현금이어도 값은 계산')
eq(z2.get('note'), '순현금 기업이라 모형 부적합(v5.4 규칙 18); ' + cl.ALTMAN_NOTE, '순현금 note + 시점/모형 note')

# ---------- RIM ----------
near(cl.rim_value(10000, [0.10, 0.10, 0.10], 0.10, 0.02, 3), 10000, 'RIM: ROE=ke 면 V=BPS0', 1e-6)
# 1년 손계산: BPS0 100, ROE 0.2, ke 0.1, g 0, payout 0
# RI1 = 0.1*100 = 10 -> PV 10/1.1 ; TV = 10*1/(0.1)/1.1 = 90.909
near(cl.rim_value(100, [0.2], 0.10, 0.0, 1), 100 + 10 / 1.1 + 10 / 0.1 / 1.1, 'RIM 1년 손계산')

# ---------- reverse DCF ----------
g_true = 0.07
ev = cl.dcf_ev(1000, g_true, 0.095, 0.02, 5)
g_hat, why = cl.solve_implied_growth(ev, 1000, 0.095, 0.02, 5)
near(g_hat, g_true, 'reverse DCF 가 g 를 회복', 1e-4)
g_bad, why = cl.solve_implied_growth(ev * 1000, 1000, 0.095, 0.02, 5)
eq((g_bad, bool(why)), (None, True), '해가 범위 밖이면 None + 사유')

ASSUME = {'ke': 0.10, 'wacc': 0.095, 'g_terminal': 0.02, 'roe_forecast': [0.12, 0.13, 0.13],
          'bps0': 10000, 'rim_years': 3, 'fcf0': 100, 'fcf_growth_years': 5,
          'net_cash': -500, 'shares': 1_000_000, 'price': 10000, 'target_price': 12000}
LA = cl.build_ledger(FS, MARKET, ASSUME)
eq(one(LA, 'rim_value_per_share')['status'], 'ok', 'assumptions 있으면 RIM ok')
near(one(LA, 'justified_pbr')['value'], (0.12 - 0.02) / (0.10 - 0.02), 'justified_pbr = ggm 식')
eq([i['value'] for i in LA['items'] if i['key'] == 'roe' and i['period'].startswith('forecast')], [0.12], '가정 ROE 가 장부 roe 항목으로')
near(one(LA, 'price_implied_roe')['value'], 0.02 + 10000 / 10000 * (0.10 - 0.02), 'price_implied_roe = g + P/BPS*(ke-g)')
# SOTP 손계산: (50*2*10 + 0.5*400 + 30 + 1000*0.5*0.8 - 500) = 1130억 / 100만주 = 113,000원; bear 는 비상장 0, 옵션 0
ASO = {**ASSUME, 'sotp': {'ebitda_half': 50, 'listed_stake': 0.5, 'listed_mcap': 400, 'unlisted_book': 30,
       'option_ev_success': 1000, 'option_discount': 0.8, 'net_debt': 500,
       'scenarios': {'base': {'multiple': 10, 'option_prob': 0.5}, 'bear': {'multiple': 10, 'option_prob': 0, 'unlisted': False}}}}
LS = cl.build_ledger(FS, MARKET, ASO)
near(one(LS, 'sotp_value_per_share_base')['value'], 1130e8 / 1_000_000, 'SOTP base 손계산')
near(one(LS, 'sotp_value_per_share_bear')['value'], (1000 + 200 - 500) * 1e8 / 1_000_000, 'SOTP bear: 비상장·옵션 0')
near(one(LS, 'sotp_upside_base')['value'], 113000 / 10000 - 1, 'SOTP upside')
ASO3 = {**ASO, 'sotp': {**ASO['sotp'], 'scenarios': {**ASO['sotp']['scenarios'], 'adj': {**ASO['sotp']['scenarios']['base'], 'adjust': 70}}}}
LS3 = cl.build_ledger(FS, MARKET, ASO3)
near(one(LS3, 'sotp_value_per_share_adj')['value'], (1130 + 70) * 1e8 / 1_000_000, 'SOTP 시나리오 조정(adjust) 가산')
# PER 시나리오: eps x per, 확률가중 (assumptions.per_scenarios / weights)
APS = {**ASSUME, 'per_scenarios': {'bear': {'eps': 1000, 'per': 5}, 'base': {'eps': 1000, 'per': 8}, 'bull': {'eps': 1200, 'per': 10}},
       'weights': {'bear': 0.3, 'base': 0.5, 'bull': 0.2}}
LP = cl.build_ledger(FS, MARKET, APS)
near(one(LP, 'per_value_base')['value'], 8000, 'PER base = eps x per')
near(one(LP, 'per_upside_bull')['value'], 12000 / 10000 - 1, 'PER bull upside')
near(one(LP, 'per_expected_value')['value'], 0.3 * 5000 + 0.5 * 8000 + 0.2 * 12000, 'PER 확률가중')
eq(one(cl.build_ledger(FS, MARKET, {**ASSUME, 'per_scenarios': {'x': {'eps': None, 'per': 5}}}), 'per_value_x')['status'] != 'ok', True, 'eps 없으면 실패 사유')
# 시나리오별 EBITDA override + 확률가중
ASO2 = {**ASO, 'sotp': {**ASO['sotp'], 'scenarios': {**ASO['sotp']['scenarios'], 'ttm': {'multiple': 10, 'option_prob': 0, 'unlisted': False, 'ebitda_annual': 60}},
        'weights': {'base': 0.5, 'ttm': 0.5}}}
LS2 = cl.build_ledger(FS, MARKET, ASO2)
near(one(LS2, 'sotp_value_per_share_ttm')['value'], (60 * 10 + 200 - 500) * 1e8 / 1_000_000, 'SOTP ebitda_annual override')
near(one(LS2, 'sotp_expected_value')['value'], 0.5 * 113000 + 0.5 * 30000, 'SOTP 확률가중 기대가치')
# 내재 확률: 시총 100억 - (1000+200+30-500=730억) = -630억 / (1000*0.8) -> 음수도 그대로 기록
near(one(LS, 'sotp_implied_option_prob')['value'], (10000 * 1_000_000 / 1e8 - 730) / 800, 'SOTP 내재 옵션 확률')
rd = one(LA, 'reverse_dcf_implied_growth')
eq(rd['status'], 'ok', 'reverse DCF ok')
# 시총 = 10000원 x 100만주 = 100억, 순현금 -500(=순부채 500) -> EV 600억
near(cl.dcf_ev(100, rd['value'], 0.095, 0.02, 5), 600, 'reverse DCF: EV = 시총(억) + 순부채 를 재현', 1e-3)
LN = cl.build_ledger(FS, MARKET, {**ASSUME, 'net_cash': 200})
eq(one(LN, 'reverse_dcf_implied_growth')['status'], 'skipped', 'EV 음수(순현금 > 시총)면 skipped')
near(one(LA, 'upside_if_view')['value'], 0.2, 'upside_if_view')

# reverse DCF 검산 허용오차 명시
res = one(LA, 'reverse_dcf_ev_residual')
eq(abs(res['value']) < cl.EV_RESIDUAL_TOL, True, f"EV 잔차 {res['value']} < 허용오차")
eq(res['tolerance_억'], cl.EV_RESIDUAL_TOL, '잔차 행에 허용오차 필드')
eq('|잔차| < 1e-06억원' in rd['note'] and '|잔차| < 1e-06억원' in res['note'], True, 'note 에 허용오차를 숫자로 적는다')

# payout 기본값 출처 / RIM 정의
rim = one(LA, 'rim_value_per_share')
eq(rim['inputs']['payout']['source'], '기본값 0 (assumptions 에 없음)', 'payout 없으면 기본값 출처')
eq(one(cl.build_ledger(FS, MARKET, {**ASSUME, 'payout': 0.2}), 'rim_value_per_share')['inputs']['payout']['source'],
   'ta/assumptions.json:payout', 'payout 있으면 assumptions 출처')
eq(rim['definition'], cl.RIM_DEFINITION, 'RIM definition')
eq('(1+g) / (ke - g) / (1+ke)^N' in cl.RIM_DEFINITION, True, 'RIM 잔여가치 식 명시')

# net_cash 부호/크기 점검 (FS 2025 net_debt = 100억)
def nc_warn(v):
    return [w for w in cl.build_ledger(FS, MARKET, {**ASSUME, 'net_cash': v})['warnings'] if w.startswith('net_cash')]
eq(nc_warn(-100), [], 'net_cash = -net_debt 면 경고 없음')
eq(nc_warn(-110), [], '크기 차이 20% 이내면 경고 없음')
eq(len(nc_warn(0)) == 1 and '템플릿 기본값' in nc_warn(0)[0], True, 'net_cash 0 인데 net_debt 있으면 경고')
eq(len(nc_warn(100)) == 1 and '부호 불일치' in nc_warn(100)[0], True, '같은 부호면 경고')
eq(len(nc_warn(-500)) == 1 and '크기 차이' in nc_warn(-500)[0], True, '크기 20% 초과 차이면 경고')
eq([w for w in L['warnings'] if w.startswith('net_cash')], [], 'assumptions 없으면 점검 안 함')

# ---------- xlsx ----------
with tempfile.TemporaryDirectory() as td:
    path = os.path.join(td, 'model_테스트.xlsx')
    cl.write_xlsx(LA, path)
    eq(os.path.exists(path), True, 'xlsx 생성')
    wb = openpyxl.load_workbook(path)
    eq(set(['Inputs', 'Calc']) <= set(wb.sheetnames), True, 'Inputs/Calc 시트')
    calc = wb['Calc']
    oks = [i for i in LA['items'] if i['status'] == 'ok']
    for it in oks:
        cell = calc[it['cell'].split('!')[1]]
        eq(isinstance(cell.value, str) and cell.value.startswith('='), True, f"{it['key']} 수식 셀 '=' 시작")
        eq(cell.value, it['excel'], f"{it['key']} ledger excel == 시트 수식")
        for ref in re.findall(r'Inputs!\$?([A-Z]+)\$?(\d+)', it['excel']):
            eq(wb['Inputs'][f'{ref[0]}{ref[1]}'].value is not None, True,
               f"{it['key']} 참조 셀 Inputs!{ref[0]}{ref[1]} 존재")
    # 수식을 직접 평가해 파이썬 값과 같은지 (엑셀 없이 차이 열 0 을 흉내).
    # eval 은 테스트 전용 -- 대상은 이 모듈이 만든 수식에 숫자만 치환한 문자열이다
    def _ev(expr, depth=0):
        def sub(m):
            sheet, row = m.group(1), int(m.group(2))
            v = (wb['Inputs'].cell(row=row, column=3).value if sheet == 'Inputs'
                 else _ev(calc.cell(row=row, column=3).value, depth + 1))
            return f'({v!r})'
        return eval(re.sub(r'(Inputs|Calc)!C(\d+)', sub, expr.lstrip('=')).replace('^', '**'))
    for it in oks:
        near(_ev(it['excel']), it['value'], f"{it['key']} 엑셀 수식 평가값 == 파이썬 값", 1e-6)
    diff_col = calc.cell(row=int(oks[0]['cell'].split('!')[1][1:]), column=5).value
    eq(str(diff_col).startswith('='), True, '차이 열도 수식')

# ---------- check ----------
LEDGER_C = {'items': [
    {'key': 'dol', 'value': 2.31, 'unit': 'x', 'period': 'FY2025 vs FY2024', 'status': 'ok'},
    {'key': 'roe', 'value': 0.148, 'unit': 'ratio', 'period': 'FY2025', 'status': 'ok', 'definition': 'ROE = X'},
]}
res = cl.check_sections({'s06': '영업레버리지(DOL) 2.3배로 높다. ROE 15.2%를 기록했다. PER 12.3배다.'},
                        LEDGER_C)
by = {r['name']: r for r in res}
eq(by.get('영업레버리지', by.get('DOL', {})).get('result'), 'PASS', 'DOL 2.3배 vs 2.31 PASS')
eq(len([r for r in res if r['text_value'] == 2.3]), 1, '같은 숫자는 한 번만 판정')
eq(by.get('ROE', {}).get('result'), 'FAIL', 'ROE 15.2% vs 0.148 FAIL')
eq(by['ROE'].get('definitions'), ['ROE = X'], 'FAIL 이면 장부 definition 을 함께 낸다')
eq(any(r['name'] == 'PER' for r in res), False, '사전에 없는 지표 무시')
res2 = cl.check_sections({'s': 'ROE는 2025년 14.8%다. 부채비율 110%.'}, LEDGER_C)
eq([(r['name'], r['result']) for r in res2], [('ROE', 'PASS'), ('부채비율', 'WARN')],
   '연도 숫자는 건너뛰고, 장부에 없는 지표는 WARN')

# 이름 공백 변형, 기간 표기 건너뛰기, VaR 신뢰수준
LEDGER_V = {'items': LEDGER_C['items'] + [
    {'key': 'justified_pbr', 'value': 1.25, 'unit': 'x', 'period': 'yr1', 'status': 'ok'},
    {'key': 'reverse_dcf_implied_growth', 'value': 0.071, 'unit': 'ratio', 'period': '5y', 'status': 'ok'},
    {'key': 'var95_1d', 'value': -0.0904, 'unit': 'ratio', 'period': 'asof', 'status': 'ok'},
    {'key': 'var95_20d', 'value': -0.21, 'unit': 'ratio', 'period': 'asof', 'status': 'ok'},
]}
rv = cl.check_sections({'s': '적정PBR 1.3배. 내재성장률 7.1%. 영업 레버리지 2.3배.'}, LEDGER_V)
eq([(r['name'], r['metric'], r['result']) for r in rv],
   [('적정PBR', '적정 PBR', 'PASS'), ('내재성장률', '내재 성장률', 'PASS'), ('영업 레버리지', '영업 레버리지', 'PASS')],
   '공백 없는/있는 이름 변형 모두 검사')
rp = cl.check_sections({'s': 'ROE 5개년 평균 14.8%로 높다.'}, LEDGER_V)
eq([(r['text_value'], r['result']) for r in rp], [(14.8, 'PASS')], '5개년 의 5 는 건너뛴다')
rq = cl.check_sections({'s': 'ROE 3분기 누적 15.2%'}, LEDGER_V)
eq([(r['text_value'], r['result']) for r in rq], [(15.2, 'WARN')], '3분기 의 3 도 건너뛰고, 분기 표기라 FY 장부값과 다르면 WARN')
r99 = cl.check_sections({'s': 'VaR(99%) -9.0% 수준이다.'}, LEDGER_V)
eq([(r['result'], r['reason'], r['var_level']) for r in r99],
   [('WARN', '장부는 95% VaR 만 있음 (본문 99%)', '99')], '99% VaR 는 95% 장부값과 비교하지 않음')
r95 = cl.check_sections({'s': 'VaR(95%) -9.0%'}, LEDGER_V)
eq([r['result'] for r in r95], ['PASS'], '95% VaR 는 비교')
r20 = cl.check_sections({'s': '20일 VaR 95% 기준 -21%'}, LEDGER_V)
eq([r['result'] for r in r20], ['PASS'], 'VaR 뒤 신뢰수준 건너뛰고 비교')
rd = cl.check_sections({'s': 'VaR 20일 -21%'}, LEDGER_V)
eq([(r['text_value'], r['result']) for r in rd], [(-21.0, 'PASS')], 'VaR 20일 의 20 은 기간')
# 보유기간별로 해당 장부값만 비교 (1d -9.04%, 20d -21%)
hv = lambda t: [(r.get('var_horizon_days'), r['result']) for r in cl.check_sections({'s': t}, LEDGER_V)]  # noqa: E731
eq(hv('VaR 20일 -9.0%'), [(20, 'FAIL')], '20일 VaR 에 1일 값을 쓰면 FAIL (1d 와 비교하지 않음)')
eq(hv('20거래일 VaR -21%'), [(20, 'PASS')], '20거래일 -> var95_20d')
eq(hv('VaR 20영업일 -21%'), [(20, 'PASS')], '20영업일 -> var95_20d')
eq(hv('VaR -21%'), [(1, 'FAIL')], '기간 표기 없으면 1일 값과만 비교')
eq(hv('일간 VaR -9.0%'), [(1, 'PASS')], '일간 -> var95_1d')
eq(hv('1일 VaR -9.0%'), [(1, 'PASS')], '1일 -> var95_1d')
eq(hv('9월 20일 기준 VaR -9.0%'), [(1, 'PASS')], '날짜(9월 20일)는 보유기간으로 읽지 않음')
r10 = cl.check_sections({'s': 'VaR 10일 -15%'}, LEDGER_V)
eq([(r['result'], r['reason']) for r in r10], [('WARN', '장부에 해당 기간 VaR 없음 (본문 10일, 장부는 1일/20일)')],
   '다른 기간은 WARN')

# 경계·단위·기준선·범위 (CJ프레시웨이 실측: 54건 중 진짜 결함은 2건, 나머지는 반기·부문·기준선·타사 수치였다)
ck = lambda t, L=LEDGER_V: [(r['metric'], r['text_value'], r['result']) for r in cl.check_sections({'s': t}, L)]  # noqa: E731
eq(ck('순부채비율 220%다.'), [], '순부채비율은 부채비율이 아니다 (이름 앞 경계)')
eq(ck('사채관리계약은 부채비율 500% 이하 유지가 조건이다.', LEDGER_C), [], '500% 이하 는 기준선이지 값이 아니다')
eq(ck('급식 영업이익률 5%대 복귀'), [], '5%대 도 기준선')
eq(ck('ROE를 그대로 넣었고 BPS에는 신종자본증권 600억원이 있다.'), [], '600억 은 지표값이 아니다')
eq(ck('ROE가 낮다. 유통 75%인 회사다.'), [], '문장이 끝나면 다음 문장의 숫자는 보지 않는다')
eq(ck('| ROE | 14.8% |'), [('ROE', 14.8, 'PASS')], '표: 이름만 있는 셀은 다음 셀의 값을 본다')
eq(ck("| '부채비율 300% 돌파' 분석 | +7.3% |", LEDGER_C), [], '표: 셀 안에 다른 글이 있으면 다음 셀로 넘어가지 않는다')
eq(ck('ROE ' + '가' * 27 + ' 11.1%'), [('ROE', 11.1, 'FAIL')], '30자 경계에서 숫자가 잘리지 않는다 (11 이 아니라 11.1)')
eq(ck('VaR(99%) -9.0% 수준이다.')[0][2], 'WARN', '"-9.0% 수준" 은 값이다 (수준은 기준선이 아님)')
eq([(r['result'], r['reason']) for r in cl.check_sections({'s': '상반기 ROE 5.7%'}, LEDGER_V)],
   [('WARN', '장부 밖 기간·범위 표기 (작성자 출처 확인)')], '반기·부문·타사 표기는 FAIL 대신 WARN')
eq([r['result'] for r in cl.check_sections({'s': '2022년 ROE 12.9%'}, LEDGER_V)], ['WARN'], '장부에 없는 연도는 WARN')
eq([r['result'] for r in cl.check_sections({'s': '2025년 ROE 12.9%'}, LEDGER_V)], ['FAIL'], '장부에 있는 연도의 다른 값은 FAIL')

# 규칙 9 -- 이익·배수 기간 정합 (하이브 critic: FY27E EPS x 12M 선행 배수)
pm = lambda ps: [(r['pos'], r['result']) for r in cl.check_period_match({'per_scenarios': ps})]  # noqa: E731
eq(pm({'a': {'eps': 1, 'per': 2}}), [('a', 'WARN')], 'basis 미표기는 WARN')
eq(pm({'a': {'eps': 1, 'per': 2, 'eps_basis': 'FY27E', 'per_basis': '2027년 컨센 PER'}}), [('a', 'PASS')], '같은 연도는 PASS')
eq(pm({'a': {'eps': 1, 'per': 2, 'eps_basis': '12M 선행', 'per_basis': 'KIS 롤링 12개월'}}), [('a', 'PASS')], '12M 끼리 PASS')
eq(pm({'a': {'eps': 1, 'per': 2, 'eps_basis': 'FY27E', 'per_basis': '12M 선행'}}), [('a', 'FAIL')], 'FY27 x 12M 은 note 없으면 FAIL')
eq(pm({'a': {'eps': 1, 'per': 2, 'eps_basis': 'FY27E', 'per_basis': '12M 선행', 'period_note': '정점 뒤 정상화 연도'}}),
   [('a', 'PASS')], 'note 가 있으면 PASS (이유를 본문에 적었다는 뜻)')
eq(pm({'a': {'eps': 1, 'per': 2, 'eps_basis': '후행 12M', 'per_basis': 'TTM'}}), [('a', 'PASS')], '후행 끼리 PASS')
eq(cl._period_class('바닥 PER'), 'unknown', '연도도 12M 도 없으면 unknown')

print('=' * 66)
if _failed:
    for label, want, got in _failed:
        print(f'  [FAIL] {label}\n      기대: {want!r}\n      실제: {got!r}')
print(f'  ta_calc_ledger 테스트: {_passed}개 통과 / {len(_failed)}개 실패')
print('=' * 66)
sys.exit(1 if _failed else 0)
