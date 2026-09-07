# -*- coding: utf-8 -*-
"""OCI홀딩스 주가 History Appendix
   (a) 메인 report_OCI홀딩스.docx 뒤에 append  (b) 별도 단독 docx 생성.
   ※ append는 generate_word_wf.py 재실행(기본 docx 재생성) 직후에 실행."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

FONT = 'Malgun Gothic'
NAVY = RGBColor(0x1E, 0x3A, 0x6D)
TPL_BLUE = RGBColor(0x26, 0x83, 0xC6)
GREY = RGBColor(0x60, 0x60, 0x60)
MAIN = 'output/OCI홀딩스/report_OCI홀딩스.docx'
STANDALONE = 'output/OCI홀딩스/OCI홀딩스_주가History_Appendix.docx'
CHART_DIR = Path('data/OCI홀딩스')
MARKER = 'Appendix. OCI홀딩스 주가 History'

CHARTS = [
 ('_hist_5y_detail', '도표 1. OCI홀딩스 주가 사이클 & 변동 이벤트 (2021~2026)',
  '빨강 원 = 상승 계기, 파랑 원 = 하락 계기'),
 ('_hist_driver_table', '도표 2. 기간별 주가 변동 요인 (Up/Down, 2020~2026)',
  '[공시] = 회사 공시 사항, [정책]·[시장]·[보도] = 외부 요인'),
 ('_hist_momentum', '도표 2. Forward 주가 모멘텀 (History → Upside, ~2027.12)',
  '● 과거 실현 이벤트, ○ 향후 기대 촉매 · Section 232 결과가 방향을 가른다'),
]


def _run(run, size=10.5, bold=False, color=None):
    run.font.size = Pt(size); run.bold = bold; run.font.name = FONT
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {}); rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), FONT)


def para(doc, text='', size=10.5, bold=False, color=None, before=0, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if text:
        _run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def chart(doc, name, caption, source):
    png = CHART_DIR / f'{name}.png'
    if not png.exists():
        print('  [경고] 이미지 없음:', png); return
    cap = para(doc, caption, size=10, bold=True, color=TPL_BLUE, before=8, after=2)
    cap.paragraph_format.keep_with_next = True
    pic_p = doc.add_paragraph(); pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.keep_with_next = True
    pic_p.add_run().add_picture(str(png), width=Cm(15.5))
    para(doc, source, size=8, color=GREY, after=8)


def build_appendix(doc, standalone=False):
    if not standalone:
        doc.add_page_break()
    para(doc, 'Appendix. OCI홀딩스 주가 History & Forward', size=15, bold=True, color=NAVY, after=5)
    para(doc,
         'OCI홀딩스 주가는 상장 이후 폴리실리콘 사이클을 따라 여러 차례 큰 폭의 상승·하락을 반복해왔다. 아래 도표는 그 '
         '사이클(적색=상승 국면, 청색=하락 국면)과 각 국면의 등락 원인을 정리한 것이다.', after=7)

    chart(doc, *CHARTS[0])

    doc.add_page_break()
    para(doc, '과거 등락의 원인 -- 무엇이 올리고 무엇이 내렸나', size=11.5, bold=True, color=NAVY, before=0, after=4)
    para(doc,
         '주가는 폴리실리콘 가격(ASP)과 미국 정책, 테마가 겹치며 크게 출렁였다. 2021~22년 상반기엔 폴리가 $39까지 오르고 '
         '사상 최대 실적(2021년 영업이익 6,260억)에 상승했고, 2022년 하반기엔 폴리 공급계약 해지·포항공장 태풍 침수·가격 '
         '급락으로 반토막 났다. 2023년 인적분할 재평가로 반등했으나, 2023 중반~2024년 중국 대증설로 폴리가 $6.55까지 '
         '무너지며 적자·저점(5.5만원)에 몰렸다. 2025년 하반기 중국 공급개혁(감산)으로 폴리가 반등하며 흑자전환 기대가 '
         '살아났고, 2026년 초 SpaceX 공급 보도와 우주·태양광 테마로 넉 달 만에 +257% 폭등해 38.8만원을 찍었다. 그러나 '
         '실적 뒷받침 없이 Section 232 발표가 지연되고 중대재해까지 겹치며 6월 -50% 급락, 현재는 21만원대 조정이다. '
         '정리하면 OCI는 폴리·정책·테마가 겹칠 때 오르고 반대에 빠지는 사이클주이며, 이번 급락은 테마 거품이 빠진 것이지 '
         '비중국 폴리 독점·미국 수직계열화라는 구조가 훼손된 것은 아니다. 거품이 걷힌 현 조정 구간은 매수 관점의 진입 '
         '기회로 판단한다.', after=6)

    para(doc, 'Forward -- 앞으로의 상승 기대', size=11.5, bold=True, color=NAVY, before=4, after=4)
    para(doc,
         '향후 주가 상승 기대는 세 갈래로 쌓인다. 단기로는 Section 232 폴리 관세가 확정되면 비중국 프리미엄이 재부각되며 '
         '재평가가 시작되고, 2분기 폴리 가동률 정상화로 흑자전환이 가시화된다. 중기로는 웨이퍼(2.7→5.4GW)·미국 텍사스 '
         'Mission Solar 셀(1→2GW)·도쿠야마 합작 OTSM 반도체 폴리(2027 준공)로 이익 기반이 폴리 한 축에서 밸류체인 전반으로 '
         '확장된다. 한화솔루션 10년 장기공급계약(1.45조원)이 기저 수요를, DCRE 분양·주주환원이 하방을 받친다. 이 촉매들이 '
         '순차 실현되면 주가는 단계적으로 상승할 것으로 기대되며, 이것이 현 구간을 매수로 보는 근거다.', after=5)

    chart(doc, *CHARTS[2])


# ── (a) 메인 docx append ──
doc = Document(MAIN)
if any(MARKER in (p.text or '') for p in doc.paragraphs):
    print('  [중단] 메인 docx에 이미 Appendix 존재. generate_word_wf로 기본 docx 먼저 재생성 필요.')
else:
    build_appendix(doc, standalone=False)
    doc.save(MAIN)
    print('  완료: 메인 docx append ->', MAIN)

# ── (b) 단독 docx ──
sd = Document()
sd.styles['Normal'].font.name = FONT
sd.styles['Normal'].font.size = Pt(10.5)
build_appendix(sd, standalone=True)
sd.save(STANDALONE)
print('  완료: 단독 docx ->', STANDALONE)
