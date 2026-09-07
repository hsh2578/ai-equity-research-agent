"""
위닝펀드 18-1 YG v6:
1. Introduction 핵심 메시지를 밸류·투자포인트에 통합
2. Introduction 슬라이드 (현재 Slide 16) 제거
3. Index 페이지에서 Introduction 라인 (shape[2], shape[13]) 비우기

통합 위치:
- Slide 8 (투자포인트 1, "빅뱅 20주년")의 박스 2 마지막 단락에 영수증 메시지
- Slide 14 (밸류에이션)의 박스 2 마지막 단락에 "컨센 +64% 위" 메시지
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pathlib import Path

SRC = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_기업분석리포트.pptx')


def replace_para_text(para, new_text):
    if not para.runs:
        return False
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ''
    return True


def delete_slide(prs, slide_idx):
    """슬라이드 완전 제거"""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[slide_idx].rId
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[slide_idx])


def main():
    p = Presentation(str(SRC))
    n_before = len(p.slides)
    print(f'Before: {n_before} slides')

    # ============ 1. 핵심 메시지 통합 ============
    # Slide 8 (idx 7, 투자포인트 1) 박스 2 (shape[2]) 마지막 단락 보강
    # 기존 마지막 단락에 영수증 메시지 추가
    s8 = p.slides[7]
    s8_box2 = list(s8.shapes)[2]
    paras = s8_box2.text_frame.paragraphs
    # 마지막 단락 텍스트 가져오기
    if len(paras) >= 1:
        last_idx = len(paras) - 1
        last_text = ''.join(r.text for r in paras[last_idx].runs)
        new_last = (
            last_text.rstrip() +
            " 좋은 회사를 실망 국면의 저점에서 사되, 8월 빅뱅과 9월 신인이라는 영수증으로 확인한다."
        )
        replace_para_text(paras[last_idx], new_last)
        print(f'[Slide 8 Box2] 영수증 메시지 추가 (last para len {len(last_text)} -> {len(new_last)})')

    # Slide 14 (idx 13, 밸류에이션) 박스 2 마지막 단락 보강
    s14 = p.slides[13]
    s14_box2 = list(s14.shapes)[2]
    paras = s14_box2.text_frame.paragraphs
    if len(paras) >= 1:
        last_idx = len(paras) - 1
        last_text = ''.join(r.text for r in paras[last_idx].runs)
        # "투자의견 'BUY'를 제시한다." 앞에 통합 메시지 끼우기
        # 마지막 단락이 보통 BUY로 끝나므로 그 앞에 삽입
        if "BUY" in last_text:
            # BUY 앞에 메시지 삽입
            new_last = last_text.replace(
                "투자의견 'BUY'를 제시한다.",
                "19개 증권사가 1Q26 실망 후 평균 목표가를 87,900원에서 78,789원으로 10% 깎았지만, 그 깎인 컨센서스조차 현재가 +64% 위에 있다는 사실이 본 리포트가 BUY를 제시하는 마지막 정량 근거다. 투자의견 'BUY'를 제시한다."
            )
        else:
            new_last = last_text.rstrip() + " 19개 증권사 컨센이 1Q26 실망으로 87,900→78,789원으로 깎였지만, 그 깎인 값조차 현재가 +64% 위다."
        replace_para_text(paras[last_idx], new_last)
        print(f'[Slide 14 Box2] 컨센 +64% 메시지 추가 (last para len {len(last_text)} -> {len(new_last)})')

    # ============ 2. Index 페이지 Introduction 라인 비우기 ============
    # 현재 Index의 shape 매핑:
    # shape[2] = '0. Introduction'
    # shape[13] = '16' (Introduction 페이지 번호)
    idx_slide = p.slides[1]
    idx_shapes = list(idx_slide.shapes)
    for sh_idx in (2, 13):
        if sh_idx >= len(idx_shapes):
            continue
        sh = idx_shapes[sh_idx]
        if sh.has_text_frame and len(sh.text_frame.paragraphs) > 0:
            old = ''.join(r.text for r in sh.text_frame.paragraphs[0].runs).strip()
            replace_para_text(sh.text_frame.paragraphs[0], '')
            print(f'[Index shape{sh_idx}] cleared (was: "{old}")')

    # ============ 3. Introduction 슬라이드 제거 ============
    # Introduction은 현재 마지막(idx 15 = Slide 16)
    delete_slide(p, n_before - 1)
    print(f'[Slides] Removed Introduction (was Slide {n_before})')

    p.save(str(SRC))
    n_after = len(p.slides)
    print(f'\nSaved: {SRC.name}')
    print(f'After: {n_after} slides')


if __name__ == '__main__':
    main()
