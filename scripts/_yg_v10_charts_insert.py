"""v10 도표 삽입 — 본문 길이별 동적 도표 영역 (양식 H=3.26 패턴 따름)"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pathlib import Path
from PIL import Image

NAVY = RGBColor(0x1E, 0x3A, 0x6D)
RED = RGBColor(0xC8, 0x10, 0x2E)
GREY = RGBColor(0x88, 0x88, 0x88)

DST = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_기업분석리포트_v10.pptx')
CHARTS = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\_charts')

CHART_L, CHART_W = 0.42, 7.42

# (PPT idx 0-based, body_chars, png_name, 도표 번호, 제목, 출처)
CHART_MAP = [
    # idx 2~5 산업
    (2, 1054, 'industry_consensus_vs_price', '도표 01.', '엔터 4사 영업이익 컨센 +6% vs 주가 -39%', '키움증권, NH투자증권, 위닝펀드'),
    (3, 1114, 'industry_ifpi', '도표 02.', 'IFPI 글로벌 음악산업 매출 11년 시계열', 'IFPI Global Music Report 2025, 위닝펀드'),
    (4, 1286, 'industry_triggers', '도표 03.', '회복 트리거 4가지 — 단기 vs 장기 카드', '키움증권 산업 리포트(2026.05.07), 위닝펀드'),
    (5, 779, 'industry_4comp', '도표 04.', 'K-POP 4사 12MF PER + 시총 비교', 'FnGuide, WiseReport, 위닝펀드'),
    # idx 6~14 기업
    (6, 1179, 'analyst_iM_p2', '도표 05.', 'YG 5년 사업·실적 구조 (iM증권 차용)', 'iM증권 기업분석(2026.04.15), 위닝펀드'),
    (7, 557, 'company_revenue_5y', '도표 06.', 'YG 매출 부문별 5년 시계열 (공연 7배 폭증)', 'FnGuide, DART 사업보고서, 위닝펀드'),
    (8, 591, 'company_comeback_cycle', '도표 07.', 'YG IP 평균 컴백 간격 — 베몬 5~6배 빠름', 'DART 사업보고서, 위닝펀드'),
    (9, 1365, 'baemon_youtube', '도표 08.', '베이비몬스터 YT 1,200만 — K팝 걸그룹 최단', 'iM증권(2026.04.15), 위닝펀드'),
    (10, 1317, 'm_first5_albums', '도표 09.', 'YG IP 첫 5개 앨범 소요 — 베몬 25개월 vs 블핑 78개월', 'DART 사업보고서, 위닝펀드'),
    (11, 1153, 'm45_4comp_gen_ip', '도표 10.', 'K-POP 4사 세대별 IP 분산 — YG 6종 가장 두꺼움', 'WiseReport, 분기보고서, 위닝펀드'),
    (12, 2723, 'm30_5y_events', '도표 11.', '5년 주가 9개 이벤트 마킹 (IP 사이클)', 'FDR, 공시, 위닝펀드'),
    (13, 435, 'ip_timeline', '도표 12.', '2026.08~2027.03 5종 IP 동시 가동 타임라인', 'DART, 공시 일정, 위닝펀드'),
    (14, 2170, 'pattern_190', '도표 13.', '2024.10~2025.08 +190% 패턴 (FDR 일봉 실측)', 'FinanceDataReader, 위닝펀드'),
    # idx 15~22 투자
    (15, 1371, 'm38_bigbang_tour', '도표 14.', '빅뱅 20주년 BIGSHOW: REBORN 14개 도시 투어 일정', 'iM증권(2026.04.15), 코첼라 공식 발표, 위닝펀드'),
    (16, 432, 'consensus_2027', '도표 15.', '컨센 2027E OP 764억 — 빅뱅 풀이어 미반영', 'FnGuide 19개사 컨센, 위닝펀드'),
    (17, 897, 'm46_boygroup_6y', '도표 16.', '6년 만의 보이그룹 — 2026년 9월 5인조 데뷔', 'YG 2026 PLAN 공식 발표, 위닝펀드'),
    (18, 468, 'analyst_교보_p2', '도표 17.', '9월 신인 보이그룹·NEXT MONSTER 일정 (교보)', '교보증권 기업분석(2026.05.11), 위닝펀드'),
    (19, 490, 'cash_ratio', '도표 18.', 'K-POP 4사 순현금/시총 — YG 30.5% 최고', 'FnGuide, 분기보고서, 위닝펀드'),
    (20, 349, 'analyst_SK_p3', '도표 19.', 'YG 6종 IP 분산 효과 시각화 (SK증권)', 'SK증권 기업분석(2026.05.11), 위닝펀드'),
    (21, 1686, 'z_score_4factor', '도표 20.', '4팩터 Z-Score — 303종목 중 1등 STRONG_BUY', 'dacon-skills-dashboard, 위닝펀드'),
    (22, 1484, 'per_historical', '도표 21.', '12MF PER 15.9 = 역사적 하단 + 블핑 부상 직전', 'FnGuide, FDR, 위닝펀드'),
    # idx 23~25 리스크
    (23, 386, 'analyst_한화_p3', '도표 22.', '리스크 1·2·3 정리 (한화투자증권)', '한화투자증권(2026.04.13), 위닝펀드'),
    (24, 459, None, None, None, None),  # 영업외 도표 없음
    (25, 990, 'm62_sept_scenario', '도표 23.', '9월 신인 보이그룹 데뷔 실패 시 시나리오', '본 리포트 자체 추정, 위닝펀드'),
    # idx 26~29 재무
    (26, 822, 'q1_revenue', '도표 24.', '1Q26 부문별 매출 (총 1,471억, +46.9%)', 'WiseReport 1Q26 분기 실적, 위닝펀드'),
    (27, 312, 'supply_20d', '도표 25.', '20일 수급 — 외인·기관 매도, 개인 흡수', 'KIS API investor_trend, 위닝펀드'),
    (28, 323, 'income_5y', '도표 26.', 'YG 5년 매출·영업이익 — 2024 절벽 후 회복', 'FnGuide SVD_Finance, 위닝펀드'),
    (29, 341, 'm69_roe_dupont', '도표 27.', 'ROE 듀폰 분해 5년 — 영업외 안정화 시 두 자릿수', 'FnGuide, 본 리포트 추정, 위닝펀드'),
    # idx 30~36 밸류
    (30, 1106, 'per_waterfall', '도표 28.', 'Target PER 24배 덧셈 워터폴', '본 리포트 자체 산정, FnGuide 4사 컨센, 위닝펀드'),
    (31, 3485, None, None, None, None),  # 8단계 분해 — 본문이 매우 김 (도표 없음)
    (32, 1130, 'm79_5_assumptions', '도표 29.', '5가지 보수 가정 매트릭스 (미반영/부분반영)', '본 리포트 자체 산정, 위닝펀드'),
    (33, 1051, 'cash_bridge', '도표 30.', '순현금 Bridge 4년 — 누적 +2,607억', 'FnGuide 컨센, 본 리포트 추정, 위닝펀드'),
    (34, 631, 'm83_ip_quarterly', '도표 31.', 'IP 라이프사이클 함수 — 분기별 IP 매출 시간표', '본 리포트 자체 산정, 위닝펀드'),
    (35, 1805, 'scenario_3', '도표 32.', 'Bear/Base/Bull 3시나리오 (호라이즌 2.5년)', '본 리포트 자체 산정, 위닝펀드'),
    (36, 1029, 'm88_price_ladder', '도표 33.', '가격 사다리 — 호라이즌별 적정주가 분포', '본 리포트 자체 산정, 위닝펀드'),
]


def get_chart_layout(body_chars):
    """본문 글자 수에 따라 도표 layout 결정 (v10_build의 get_layout과 동일 매핑)"""
    if body_chars < 700:
        return dict(chart_t=3.85, chart_h=6.50, caption_t=3.55, source_t=10.55)
    elif body_chars < 1500:
        return dict(chart_t=5.55, chart_h=4.80, caption_t=5.25, source_t=10.55)
    elif body_chars < 2500:
        return dict(chart_t=6.85, chart_h=3.50, caption_t=6.55, source_t=10.55)
    else:
        return None  # 도표 없음


def insert_chart(slide, png_path, chart_no, title, source, layout):
    """캡션 + 도표 + 출처 삽입 (양식 H=3.26 패턴 기준)"""
    L, T, W, H = CHART_L, layout['chart_t'], CHART_W, layout['chart_h']

    # 캡션
    tb = slide.shapes.add_textbox(Inches(L), Inches(layout['caption_t']), Inches(W), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = chart_no
    r1.font.size = Pt(10); r1.font.bold = True; r1.font.color.rgb = RED
    r2 = p.add_run(); r2.text = f'  {title}'
    r2.font.size = Pt(10); r2.font.bold = True; r2.font.color.rgb = NAVY

    # 도표 이미지 (비율 자동 fit + 중앙)
    if png_path.exists():
        try:
            im = Image.open(str(png_path))
            iw, ih = im.size
            if W * (ih / iw) <= H:
                actual_w = W; actual_h = W * (ih / iw)
                off_l = L; off_t = T + (H - actual_h) / 2
            else:
                actual_h = H; actual_w = H * (iw / ih)
                off_l = L + (W - actual_w) / 2; off_t = T
            slide.shapes.add_picture(str(png_path), Inches(off_l), Inches(off_t),
                                     Inches(actual_w), Inches(actual_h))
        except Exception:
            slide.shapes.add_picture(str(png_path), Inches(L), Inches(T),
                                     Inches(W), Inches(H))

    # 출처
    tb2 = slide.shapes.add_textbox(Inches(L), Inches(layout['source_t']), Inches(W), Inches(0.22))
    tf2 = tb2.text_frame
    tf2.margin_left = Emu(0); tf2.margin_right = Emu(0)
    tf2.margin_top = Emu(0); tf2.margin_bottom = Emu(0)
    p2 = tf2.paragraphs[0]
    r3 = p2.add_run(); r3.text = f'자료: {source}'
    r3.font.size = Pt(8); r3.font.color.rgb = GREY; r3.font.italic = True


def remove_unused(slide):
    """양식 GROUP + 도표 영역 LINE 제거"""
    to_remove = []
    for sh in slide.shapes:
        l = Emu(sh.left).inches if sh.left else 0
        t = Emu(sh.top).inches if sh.top else 0
        w = Emu(sh.width).inches if sh.width else 0
        if sh.shape_type == 6 and l < 1.0 and w > 5.0:
            to_remove.append(sh)
        elif sh.shape_type == 9 and 3.0 < t < 11.10:
            to_remove.append(sh)
    for sh in to_remove:
        sh._element.getparent().remove(sh._element)


def main():
    prs = Presentation(str(DST))
    print(f'PPT 로드: {len(prs.slides)}장\n')

    inserted = 0
    missing = []
    skipped_long = 0
    print('=== v10 도표 삽입 (본문 길이별 동적 H) ===')
    for slide_idx, body_chars, fname, no, title, src in CHART_MAP:
        if slide_idx >= len(prs.slides):
            print(f'  Slide {slide_idx+1} out of range')
            continue
        slide = prs.slides[slide_idx]
        remove_unused(slide)

        layout = get_chart_layout(body_chars)
        if layout is None:
            print(f'  ⏭  Slide {slide_idx+1:2}: 본문 {body_chars}자 → 도표 영역 없음')
            skipped_long += 1
            continue
        if fname is None:
            print(f'  ⏭  Slide {slide_idx+1:2}: 도표 매핑 없음')
            continue

        png = CHARTS / f'{fname}.png'
        if not png.exists():
            missing.append(fname)
            print(f'  ✗  Slide {slide_idx+1:2}: PNG 누락 {fname}')
            continue

        insert_chart(slide, png, no, title, src, layout)
        inserted += 1
        print(f'  ✓  Slide {slide_idx+1:2}: H={layout["chart_h"]:.2f}  {no} {title[:30]}...')

    prs.save(str(DST))
    print(f'\n도표 삽입: {inserted}개 / 긴 본문 건너뜀: {skipped_long}개')
    if missing: print(f'⚠ 누락: {missing}')
    print(f'저장: {DST.stat().st_size / 1024:.1f} KB')


if __name__ == '__main__':
    main()
