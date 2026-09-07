"""
위닝펀드 18-1 YG pptx 도표 교체:
지정 슬라이드의 기존 큰 Picture(W>3, H>2)를 새 PNG로 같은 위치·크기로 교체.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Emu
from pathlib import Path

SRC = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_기업분석리포트.pptx')


def swap_chart(slide_num: int, new_png: Path):
    p = Presentation(str(SRC))
    slide = p.slides[slide_num - 1]
    target = None
    for shape in slide.shapes:
        if shape.shape_type != 13:
            continue
        w = Emu(shape.width).inches
        h = Emu(shape.height).inches
        if w > 3 and h > 2:
            target = shape
            break
    if not target:
        print(f'[ERR] Slide {slide_num}: no large picture found')
        return

    left, top, width, height = target.left, target.top, target.width, target.height
    print(f'[Slide {slide_num}] Existing: L={Emu(left).inches:.2f}" T={Emu(top).inches:.2f}" '
          f'W={Emu(width).inches:.2f}" H={Emu(height).inches:.2f}"')

    # 삭제: shape의 XML element를 부모에서 제거
    sp = target._element
    sp.getparent().remove(sp)

    # 새 그림 삽입
    pic = slide.shapes.add_picture(str(new_png), left, top, width, height)
    print(f'[Slide {slide_num}] Inserted: {new_png.name}')

    p.save(str(SRC))
    print(f'Saved: {SRC.name}')


if __name__ == '__main__':
    import sys
    base = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터')
    if len(sys.argv) >= 3:
        slide_num = int(sys.argv[1])
        png = base / sys.argv[2]
        swap_chart(slide_num, png)
    else:
        swap_chart(12, base / '_chart_s12_history_shocks.png')
