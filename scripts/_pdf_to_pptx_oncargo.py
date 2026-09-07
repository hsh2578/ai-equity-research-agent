"""
PDF -> PPTX 변환:
PDF 각 페이지를 고해상도 PNG로 추출 후 PPTX 슬라이드에 이미지로 삽입.
슬라이드 크기는 PDF 페이지 비율에 맞춤.
"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import fitz
from pptx import Presentation
from pptx.util import Emu, Inches
from pathlib import Path

PDF = Path(r'C:\Users\hsh\Desktop\OnCargo VC 투자제안서.pdf')
OUT_PPTX = Path(r'C:\Users\hsh\Desktop\OnCargo VC 투자제안서.pptx')
TMP_DIR = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\_oncargo_pages')
TMP_DIR.mkdir(parents=True, exist_ok=True)


def main():
    doc = fitz.open(str(PDF))
    n = len(doc)
    print(f'PDF: {n} pages')

    # 첫 페이지 크기로 PPT 크기 결정
    first = doc[0]
    w_pt = first.rect.width   # in points (1pt = 1/72")
    h_pt = first.rect.height
    w_in = w_pt / 72.0
    h_in = h_pt / 72.0
    print(f'Page size: {w_pt:.0f}x{h_pt:.0f} pt = {w_in:.2f}x{h_in:.2f} inch')

    # PPTX 생성
    prs = Presentation()
    prs.slide_width = Inches(w_in)
    prs.slide_height = Inches(h_in)
    blank = prs.slide_layouts[6]  # Blank layout

    # 페이지별 PNG 추출 + 슬라이드 추가
    zoom = 2.0  # 2배 해상도 (PDF 72dpi -> 144dpi)
    mat = fitz.Matrix(zoom, zoom)
    for i in range(n):
        page = doc[i]
        pix = page.get_pixmap(matrix=mat)
        png_path = TMP_DIR / f'page_{i+1:02d}.png'
        pix.save(str(png_path))

        # 새 슬라이드
        slide = prs.slides.add_slide(blank)
        # 이미지 풀사이즈로 삽입
        slide.shapes.add_picture(str(png_path), 0, 0,
                                  width=prs.slide_width, height=prs.slide_height)
        print(f'  Slide {i+1}: {png_path.name}')

    prs.save(str(OUT_PPTX))
    print(f'\nSaved: {OUT_PPTX}')
    print(f'Size: {os.path.getsize(OUT_PPTX) / 1024:.1f} KB')


if __name__ == '__main__':
    main()
