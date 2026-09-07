"""차트 PNG를 PPT 도표 영역에 삽입 + 도표 번호·출처 표기"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pathlib import Path

NAVY = RGBColor(0x1E, 0x3A, 0x6D)
RED = RGBColor(0xC8, 0x10, 0x2E)
GREY = RGBColor(0x88, 0x88, 0x88)

DST = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_기업분석리포트_v8.pptx')
CHARTS = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\_charts')

# 슬라이드별 도표 매핑 (idx, [상단 도표, 하단 도표])
# 각 도표: (파일명, 도표 번호, 도표 제목, 출처)
CHART_MAP = {
    2: [  # Slide 3 산업분석 1
        ('industry_consensus_vs_price', '도표 01.', '엔터 4사 영업이익 컨센 +6% vs 주가 -39%', '키움증권, NH투자증권, 위닝펀드'),
        ('industry_ifpi', '도표 02.', 'IFPI 글로벌 음악산업 매출 11년 시계열', 'IFPI Global Music Report 2025, 위닝펀드'),
    ],
    3: [  # Slide 4 산업분석 2
        ('industry_triggers', '도표 03.', '회복 트리거 4가지 — 단기 vs 장기 카드', '키움증권 산업 리포트(2026.05.07), 위닝펀드'),
        ('industry_4comp', '도표 04.', 'K-POP 4사 12MF PER + 시총 비교', 'FnGuide, WiseReport, 위닝펀드'),
    ],
    4: [  # Slide 5 기업분석 1
        ('company_revenue_5y', '도표 05.', 'YG 매출 부문별 5년 시계열 (공연 7배 폭증)', 'FnGuide, DART 사업보고서, 위닝펀드'),
        ('analyst_iM_p2', '도표 06.', 'YG 5년 사업·실적 구조 (iM증권 차용)', 'iM증권 기업분석(2026.04.15), 위닝펀드'),
    ],
    5: [  # Slide 6 기업분석 2
        ('company_comeback_cycle', '도표 07.', 'YG IP 평균 컴백 간격 — 베몬 5~6배 빠름', 'DART 사업보고서, 위닝펀드'),
        ('baemon_youtube', '도표 08.', '베이비몬스터 YT 1,200만 — K팝 걸그룹 최단', 'iM증권(2026.04.15), 위닝펀드'),
    ],
    6: [  # Slide 7 기업분석 3
        ('analyst_유진_p2', '도표 09.', '베이비몬스터 컴백·월투 일정 (유진투자증권)', '유진투자증권 기업분석(2026.05.11), 위닝펀드'),
        ('analyst_교보_p2', '도표 10.', 'YG 6종 IP 포트폴리오·신인 데뷔 일정', '교보증권 기업분석(2026.05.11), 위닝펀드'),
    ],
    7: [  # Slide 8 기업분석 4 (5년 주가)
        ('pattern_190', '도표 11.', '2024.10~2025.08 +190% 패턴 (FDR 일봉 실측)', 'FinanceDataReader, 위닝펀드'),
        ('ip_timeline', '도표 12.', '2026.08~2027.03 5종 IP 동시 가동 타임라인', 'DART 사업보고서, 공시 일정, 위닝펀드'),
    ],
    8: [  # Slide 9 투자포인트 1
        ('consensus_2027', '도표 13.', '컨센 2027E OP 764억 — 빅뱅 풀이어 미반영', 'FnGuide 19개사 컨센, 위닝펀드'),
        ('analyst_iM_p2', '도표 14.', '빅뱅 20주년 투어 매출·이익 추정 (iM)', 'iM증권 기업분석(2026.04.15), 위닝펀드'),
    ],
    9: [  # Slide 10 투자포인트 2
        ('analyst_교보_p2', '도표 15.', '9월 신인 보이그룹·NEXT MONSTER 일정 (교보)', '교보증권 기업분석(2026.05.11), 위닝펀드'),
        ('analyst_SK_p3', '도표 16.', 'YG 6종 IP 분산 효과 시각화 (SK증권)', 'SK증권 기업분석(2026.05.11), 위닝펀드'),
    ],
    10: [  # Slide 11 투자포인트 3
        ('z_score_4factor', '도표 17.', '4팩터 Z-Score — 303종목 중 1등', 'dacon-skills-dashboard, 위닝펀드'),
        ('per_historical', '도표 18.', '12MF PER 15.9 = 역사적 하단 + 블핑 부상 직전 구간', 'FnGuide, FDR, 위닝펀드'),
    ],
    11: [  # Slide 12 리스크
        ('analyst_한화_p3', '도표 19.', '리스크 1·2·3 정리 (한화투자증권)', '한화투자증권 기업분석(2026.04.13), 위닝펀드'),
        ('cash_ratio', '도표 20.', 'K-POP 4사 순현금/시총 — YG 30.5% 최고', 'FnGuide, 분기보고서, 위닝펀드'),
    ],
    12: [  # Slide 13 재무 1
        ('q1_revenue', '도표 21.', '1Q26 부문별 매출 (총 1,471억, +46.9%)', 'WiseReport 1Q26 분기 실적, 위닝펀드'),
        ('supply_20d', '도표 22.', '20일 수급 — 외인·기관 매도, 개인 흡수', 'KIS API investor_trend, 위닝펀드'),
    ],
    13: [  # Slide 14 재무 2
        ('income_5y', '도표 23.', 'YG 5년 매출·영업이익 — 2024 절벽 후 회복', 'FnGuide SVD_Finance, 위닝펀드'),
        ('cash_ratio', '도표 24.', '재무 건전성 4사 비교 (순현금/시총)', 'FnGuide, 위닝펀드'),
    ],
    14: [  # Slide 15 밸류 1
        ('eps_waterfall', '도표 25.', '2028E EPS 8단계 분해 → 3,700원', '본 리포트 자체 추정, 위닝펀드'),
        ('per_waterfall', '도표 26.', 'Target PER 24배 덧셈 워터폴', '본 리포트 자체 산정, FnGuide 4사 컨센, 위닝펀드'),
    ],
    15: [  # Slide 16 밸류 2
        ('cash_bridge', '도표 27.', '순현금 Bridge 4년 — 누적 +2,607억', 'FnGuide 컨센, 본 리포트 추정, 위닝펀드'),
        ('scenario_3', '도표 28.', 'Bear/Base/Bull 3시나리오 (호라이즌 2.5년)', '본 리포트 자체 산정, 위닝펀드'),
    ],
}


def remove_group_shapes(slide):
    """도표 placeholder GROUP + 도표 영역 침범 LINE 제거"""
    to_remove = []
    for sh in slide.shapes:
        l = Emu(sh.left).inches if sh.left else 0
        t = Emu(sh.top).inches if sh.top else 0
        w = Emu(sh.width).inches if sh.width else 0
        # GROUP 도표 placeholder (shape_type 6)
        if sh.shape_type == 6 and l < 1.0 and w > 5.0:
            to_remove.append(sh)
        # LINE shape (shape_type 9) — 도표 영역(T=1.85~5.80 또는 6.85~11.10) 침범 시 제거
        elif sh.shape_type == 9:
            if 1.85 < t < 5.80 or 6.85 < t < 11.10:
                to_remove.append(sh)
    for sh in to_remove:
        sh._element.getparent().remove(sh._element)


def insert_chart(slide, png_path, chart_no, title, source, position='top'):
    """차트 PNG 삽입 + 도표 번호·제목·출처 텍스트 박스 (v8 — 도표 영역 +12~20% 확대)"""
    from pptx.util import Inches
    if position == 'top':
        # 상단: 캡션 T=1.95 / 도표 T=2.25 H=3.30 / 출처 T=5.58
        # 본문 박스(T=0.83~1.82)와 0.13 간격, 도표 H 2.95→3.30 (+12%)
        L, T, W, H = 0.42, 2.25, 7.42, 3.30
        text_t = 1.95
        src_t = 5.58
    else:  # bottom
        # 하단: 캡션 T=6.95 / 도표 T=7.25 H=3.65 / 출처 T=10.93
        # 하단 본문(T=5.86~6.85)과 0.10 간격, 도표 H 3.05→3.65 (+20%)
        L, T, W, H = 0.42, 7.25, 7.42, 3.65
        text_t = 6.95
        src_t = 10.93

    # 도표 번호 + 제목 (좌측 빨강 보조 마크 + 텍스트)
    tb = slide.shapes.add_textbox(Inches(L), Inches(text_t), Inches(W), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = chart_no
    r1.font.size = Pt(9); r1.font.bold = True; r1.font.color.rgb = RED
    r2 = p.add_run(); r2.text = f'  {title}'
    r2.font.size = Pt(9); r2.font.bold = True; r2.font.color.rgb = NAVY

    # 차트 이미지 삽입 (가로 fit, 세로 비율 유지)
    if png_path.exists():
        from PIL import Image
        try:
            im = Image.open(str(png_path))
            iw, ih = im.size
            target_w = W
            target_h_calc = target_w * (ih / iw)
            if target_h_calc <= H:
                # 가로 기준 fit
                actual_h = target_h_calc
                actual_w = target_w
                offset_l = L
                offset_t = T + (H - actual_h) / 2  # 세로 중앙
            else:
                # 세로 기준 fit
                actual_h = H
                actual_w = H * (iw / ih)
                offset_l = L + (W - actual_w) / 2  # 가로 중앙
                offset_t = T
            slide.shapes.add_picture(str(png_path), Inches(offset_l), Inches(offset_t),
                                     Inches(actual_w), Inches(actual_h))
        except Exception:
            slide.shapes.add_picture(str(png_path), Inches(L), Inches(T), Inches(W), Inches(H))

    # 출처 (작은 텍스트)
    tb2 = slide.shapes.add_textbox(Inches(L), Inches(src_t), Inches(W), Inches(0.22))
    tf2 = tb2.text_frame
    tf2.margin_left = Emu(0); tf2.margin_right = Emu(0)
    tf2.margin_top = Emu(0); tf2.margin_bottom = Emu(0)
    p2 = tf2.paragraphs[0]
    r3 = p2.add_run(); r3.text = f'자료: {source}'
    r3.font.size = Pt(7); r3.font.color.rgb = GREY; r3.font.italic = True


def insert_cover_chart(slide):
    """Cover 좌측 하단 주가 차트"""
    from pptx.util import Inches
    png = CHARTS / 'cover_price.png'
    if not png.exists(): return
    # 양식 위치: L=0.49, T=9.86, W=2.29, H=1.34
    # 기존 차트가 있을 수 있으니 제거 후 삽입
    to_remove = []
    for sh in slide.shapes:
        if sh.shape_type == 3:  # CHART
            to_remove.append(sh)
    for sh in to_remove:
        sh._element.getparent().remove(sh._element)
    slide.shapes.add_picture(str(png), Inches(0.49), Inches(9.86), Inches(2.29), Inches(1.34))


def main():
    prs = Presentation(str(DST))
    print(f'PPT 로드: {len(prs.slides)}개\n')

    # Cover 차트
    print('=== Cover 차트 삽입 ===')
    insert_cover_chart(prs.slides[0])

    # 본문 슬라이드 도표
    print('\n=== 본문 도표 삽입 ===')
    for slide_idx, charts in CHART_MAP.items():
        if slide_idx >= len(prs.slides):
            print(f'  Slide {slide_idx+1} out of range, skip')
            continue
        slide = prs.slides[slide_idx]
        # 기존 GROUP shape (placeholder) 제거
        remove_group_shapes(slide)
        # 상단·하단 도표 삽입
        for i, (fname, no, title, src) in enumerate(charts):
            png = CHARTS / f'{fname}.png'
            position = 'top' if i == 0 else 'bottom'
            insert_chart(slide, png, no, title, src, position)
            mark = '✓' if png.exists() else '✗'
            print(f'  {mark} Slide {slide_idx+1} [{position}]: {no} {title[:30]}...')

    prs.save(str(DST))
    print(f'\n저장 완료: {DST.stat().st_size / 1024:.1f} KB')


if __name__ == '__main__':
    main()
