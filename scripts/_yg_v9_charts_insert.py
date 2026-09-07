"""v9 도표 삽입 — 1 슬라이드 = 1 도표 (큰 H=6.50)"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pathlib import Path
from PIL import Image

NAVY = RGBColor(0x1E, 0x3A, 0x6D)
RED = RGBColor(0xC8, 0x10, 0x2E)
GREY = RGBColor(0x88, 0x88, 0x88)

DST = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_기업분석리포트_v9.pptx')
CHARTS = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\_charts')

# v9 도표 좌표 (1 슬라이드 1 도표, A4 세로 11.69)
CHART_L, CHART_T, CHART_W, CHART_H = 0.42, 4.25, 7.42, 6.50
CAPTION_T = 3.95
SOURCE_T = 10.95

# Slide idx (0-based) → (PNG, 도표 번호, 제목, 출처)
CHART_MAP_V9 = {
    # Slide 3-6 산업
    2:  ('industry_consensus_vs_price', '도표 01.', '엔터 4사 영업이익 컨센 +6% vs 주가 -39%', '키움증권, NH투자증권, 위닝펀드'),
    3:  ('industry_ifpi', '도표 02.', 'IFPI 글로벌 음악산업 매출 11년 시계열', 'IFPI Global Music Report 2025, 위닝펀드'),
    4:  ('industry_triggers', '도표 03.', '회복 트리거 4가지 — 단기 vs 장기 카드', '키움증권 산업 리포트(2026.05.07), 위닝펀드'),
    5:  ('industry_4comp', '도표 04.', 'K-POP 4사 12MF PER + 시총 비교', 'FnGuide, WiseReport, 위닝펀드'),
    # Slide 7-14 기업
    6:  ('company_revenue_5y', '도표 05.', 'YG 매출 부문별 5년 시계열 (공연 7배 폭증)', 'FnGuide, DART 사업보고서, 위닝펀드'),
    7:  ('analyst_iM_p2', '도표 06.', 'YG 5년 사업·실적 구조 (iM증권 차용)', 'iM증권 기업분석(2026.04.15), 위닝펀드'),
    8:  ('company_comeback_cycle', '도표 07.', 'YG IP 평균 컴백 간격 — 베몬 5~6배 빠름', 'DART 사업보고서, 위닝펀드'),
    9:  ('baemon_youtube', '도표 08.', '베이비몬스터 YT 1,200만 — K팝 걸그룹 최단', 'iM증권(2026.04.15), 위닝펀드'),
    10: ('m_first5_albums', '도표 09.', 'YG IP 첫 5개 앨범 소요 — 베몬 25개월 vs 블핑 78개월', 'DART 사업보고서, 위닝펀드'),
    11: ('ip_timeline', '도표 10.', '2026.08~2027.03 5종 IP 동시 가동 타임라인', 'DART, 공시 일정, 위닝펀드'),
    12: ('pattern_190', '도표 11.', '2024.10~2025.08 +190% 패턴 (FDR 일봉 실측)', 'FinanceDataReader, 위닝펀드'),
    13: ('m30_5y_events', '도표 12.', '5년 주가 9개 이벤트 마킹 (IP 사이클)', 'FDR, 공시, 위닝펀드'),
    # Slide 15-20 투자
    14: ('analyst_iM_p2', '도표 13.', '빅뱅 20주년 투어 매출·이익 추정 (iM)', 'iM증권(2026.04.15), 위닝펀드'),
    15: ('consensus_2027', '도표 14.', '컨센 2027E OP 764억 — 빅뱅 풀이어 미반영', 'FnGuide 19개사 컨센, 위닝펀드'),
    16: ('analyst_교보_p2', '도표 15.', '9월 신인 보이그룹·NEXT MONSTER 일정 (교보)', '교보증권(2026.05.11), 위닝펀드'),
    17: ('m45_4comp_gen_ip', '도표 16.', 'K-POP 4사 세대별 IP 분산 — YG 가장 두꺼움', 'WiseReport, 분기보고서, 위닝펀드'),
    18: ('cash_ratio', '도표 17.', 'K-POP 4사 순현금/시총 — YG 30.5% 최고', 'FnGuide, 분기보고서, 위닝펀드'),
    19: ('per_historical', '도표 18.', '12MF PER 15.9 = 역사적 하단 + 블핑 부상 직전', 'FnGuide, FDR, 위닝펀드'),
    # Slide 21-22 리스크
    20: ('analyst_한화_p3', '도표 19.', '리스크 1·2·3 정리 (한화투자증권)', '한화투자증권(2026.04.13), 위닝펀드'),
    21: ('z_score_4factor', '도표 20.', '4팩터 Z-Score — 303종목 중 1등 STRONG_BUY', 'dacon-skills-dashboard, 위닝펀드'),
    # Slide 23-26 재무
    22: ('q1_revenue', '도표 21.', '1Q26 부문별 매출 (총 1,471억, +46.9%)', 'WiseReport 1Q26 분기 실적, 위닝펀드'),
    23: ('supply_20d', '도표 22.', '20일 수급 — 외인·기관 매도, 개인 흡수', 'KIS API investor_trend, 위닝펀드'),
    24: ('income_5y', '도표 23.', 'YG 5년 매출·영업이익 — 2024 절벽 후 회복', 'FnGuide SVD_Finance, 위닝펀드'),
    25: ('m69_roe_dupont', '도표 24.', 'ROE 듀폰 분해 5년 — 영업외 안정화 시 두 자릿수', 'FnGuide, 본 리포트 추정, 위닝펀드'),
    # Slide 27-30 밸류
    26: ('eps_waterfall', '도표 25.', '2028E EPS 8단계 분해 → 3,700원', '본 리포트 자체 추정, 위닝펀드'),
    27: ('per_waterfall', '도표 26.', 'Target PER 24배 덧셈 워터폴', '본 리포트 자체 산정, FnGuide 4사 컨센, 위닝펀드'),
    28: ('cash_bridge', '도표 27.', '순현금 Bridge 4년 — 누적 +2,607억', 'FnGuide 컨센, 본 리포트 추정, 위닝펀드'),
    29: ('scenario_3', '도표 28.', 'Bear/Base/Bull 3시나리오 (호라이즌 2.5년)', '본 리포트 자체 산정, 위닝펀드'),
}


def insert_chart_v9(slide, png_path, chart_no, title, source):
    """v9: 단일 큰 도표 삽입 (H=6.50, A4 세로 56%)"""
    # 캡션 (도표 번호 + 제목)
    tb = slide.shapes.add_textbox(Inches(CHART_L), Inches(CAPTION_T), Inches(CHART_W), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = chart_no
    r1.font.size = Pt(10); r1.font.bold = True; r1.font.color.rgb = RED
    r2 = p.add_run(); r2.text = f'  {title}'
    r2.font.size = Pt(10); r2.font.bold = True; r2.font.color.rgb = NAVY

    # 도표 이미지 (가로/세로 비율 자동 fit + 중앙 정렬)
    if png_path.exists():
        try:
            im = Image.open(str(png_path))
            iw, ih = im.size
            target_w, target_h = CHART_W, CHART_H
            if target_w * (ih / iw) <= target_h:
                actual_w = target_w
                actual_h = target_w * (ih / iw)
                offset_l = CHART_L
                offset_t = CHART_T + (target_h - actual_h) / 2
            else:
                actual_h = target_h
                actual_w = target_h * (iw / ih)
                offset_l = CHART_L + (target_w - actual_w) / 2
                offset_t = CHART_T
            slide.shapes.add_picture(str(png_path), Inches(offset_l), Inches(offset_t),
                                     Inches(actual_w), Inches(actual_h))
        except Exception as e:
            slide.shapes.add_picture(str(png_path), Inches(CHART_L), Inches(CHART_T),
                                     Inches(CHART_W), Inches(CHART_H))

    # 출처
    tb2 = slide.shapes.add_textbox(Inches(CHART_L), Inches(SOURCE_T), Inches(CHART_W), Inches(0.22))
    tf2 = tb2.text_frame
    tf2.margin_left = Emu(0); tf2.margin_right = Emu(0)
    tf2.margin_top = Emu(0); tf2.margin_bottom = Emu(0)
    p2 = tf2.paragraphs[0]
    r3 = p2.add_run(); r3.text = f'자료: {source}'
    r3.font.size = Pt(8); r3.font.color.rgb = GREY; r3.font.italic = True


def remove_chart_placeholder(slide):
    """양식 잔재 GROUP 도표 placeholder + 도표 영역 침범 LINE 제거"""
    to_remove = []
    for sh in slide.shapes:
        l = Emu(sh.left).inches if sh.left else 0
        t = Emu(sh.top).inches if sh.top else 0
        w = Emu(sh.width).inches if sh.width else 0
        if sh.shape_type == 6 and l < 1.0 and w > 5.0:
            to_remove.append(sh)
        elif sh.shape_type == 9 and 3.50 < t < 11.10:
            to_remove.append(sh)
    for sh in to_remove:
        sh._element.getparent().remove(sh._element)


def main():
    prs = Presentation(str(DST))
    print(f'PPT 로드: {len(prs.slides)}장\n')

    print('=== v9 도표 삽입 (1 슬라이드 1 도표, H=6.50) ===')
    inserted = 0
    missing = []
    for slide_idx, (fname, no, title, src) in CHART_MAP_V9.items():
        if slide_idx >= len(prs.slides):
            print(f'  Slide {slide_idx+1} out of range, skip')
            continue
        slide = prs.slides[slide_idx]
        remove_chart_placeholder(slide)
        png = CHARTS / f'{fname}.png'
        insert_chart_v9(slide, png, no, title, src)
        mark = '✓' if png.exists() else '✗'
        if not png.exists(): missing.append(fname)
        print(f'  {mark} Slide {slide_idx+1:2}: {no} {title[:35]}...')
        if png.exists(): inserted += 1

    prs.save(str(DST))
    print(f'\n도표 삽입 완료: {inserted}개')
    if missing:
        print(f'⚠ 누락 PNG: {missing}')
    print(f'저장: {DST.stat().st_size / 1024:.1f} KB')


if __name__ == '__main__':
    main()
