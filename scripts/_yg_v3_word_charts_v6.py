"""v6 — 부정확 스크린샷 모두 제거, 검증된 자체 차트 + 정확한 IR 자료만 사용"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

DOC = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_본문_편집용.docx')
CHARTS = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\_charts')

# 검증된 매핑만 (부정확 PDF 크롭 제거됨)
EXACT = [
    # --- 산업분석 (자체 차트만) ---
    ('IFPI 글로벌 음악산업 매출 11년', 'industry_ifpi', 'IFPI Global Music Report 2025, 위닝펀드'),
    ('엔터 4사 영업이익 컨센 +6% vs 주가 -39% 듀얼', 'industry_consensus_vs_price', '키움증권(2026.05.07), FnGuide, 위닝펀드'),
    ('업종별 연초 대비 수익률 비교', 'm03_sector_return', '키움증권 산업리포트(2026.05.07), FDR, 위닝펀드'),
    ('엔터 4사 매출 성장률 추이', 'm04_revenue_growth', '키움증권(2026.05.07), 위닝펀드'),
    ('글로벌 빅3 레이블 (UMG·SME·WMG) 10년', 'm05_global_big3', '키움증권, Bloomberg, 위닝펀드'),
    ('BTS vs 테일러 스위프트 월드투어 매출 구조', 'm06_bts_taylor', '키움증권 산업리포트, 위닝펀드'),
    ('슈퍼팬 비율 비교', 'm07_superfan', 'Luminate 2024 Music Report, 위닝펀드'),
    ('4개 회복 트리거 산업 구조 카드(중장기 2) vs 단기 카드(2)', 'industry_triggers', '키움증권 산업리포트(2026.05.07), 위닝펀드'),
    ('4사 비즈니스 모델 비교표', 'industry_4comp', 'FnGuide, WiseReport, 위닝펀드'),
    # --- 기업분석 (자체 + 검증된 IR) ---
    ('28년 데뷔 IP 타임라인 (지누션 1997', 'm10_28y_timeline', 'YG 사업보고서(2026.03), 위닝펀드'),
    ('매출 부문별 5년 시계열 (음반·공연·음원·기타 stacked', 'company_revenue_5y', 'FnGuide SVD_Finance, YG 사업보고서, 위닝펀드'),
    ('베이비몬스터 유튜브 구독자 추이 (2024.04 데뷔', 'baemon_youtube', 'iM증권(2026.04.15), YouTube, 위닝펀드'),
    ('와이지 IP별 데뷔 후 첫 5개 앨범 소요 기간 비교', 'm_first5_albums', '한터차트, YG 사업보고서, 위닝펀드'),
    ('와이지 IP별 평균 컴백 간격 비교', 'company_comeback_cycle', 'YG 사업보고서, 한터차트, 위닝펀드'),
    # iM증권 p2 (검증 완료 — 표1·2·3)
    ('베이비몬스터 일본 팝업스토어 + MD 라인업 (iM증권', 'analyst_iM_p2', 'iM증권 기업분석(2026.04.15) p.2, 위닝펀드'),
    # 제거: m_ir_4comp_debut — 도표 메모와 내용 매칭 부적합 (4사 데뷔 SNS 비교 vs 베몬·글로벌 걸그룹 비교)
    # 제거: m_ir_yg_album_concert — 도표 메모와 내용 매칭 부적합 (음반 판매량 vs 6종 IP 포트폴리오 도식)
    # 자체 차트
    ('5년 일봉 주가 + 9개 이벤트 마킹', 'm30_5y_events', 'FinanceDataReader 일봉, 위닝펀드'),
    ('2024.10~2025.08 +190% 상승 구간 확대', 'pattern_190', 'FinanceDataReader 일봉, 위닝펀드'),
    # --- 라인업·사이클 (자체) ---
    ('2026 하반기 5종 IP 동시 가동 일정표', 'ip_timeline', 'YG 사업보고서, 공시 일정, 위닝펀드'),
    ('24.10~25.08 vs 26.08~27.07 사이클 비교 매트릭스', 'm35_cycle_compare', '본 리포트 자체 추정, FnGuide, 위닝펀드'),
    # --- 투자포인트 (자체 + 검증 IR) ---
    ('빅뱅 BIGSHOW: REBORN 14개 도시 투어 일정표', 'm38_bigbang_tour', 'YG 공식 발표(2026.04), 위닝펀드'),
    # m_im_bigbang_coachella는 검증된 IR (코첼라 + MD + 블핑 + DEADLINE 4종)
    ('빅뱅 코첼라 4월 12·19일 무대 1주차 영상 1,750만뷰', 'm_im_bigbang_coachella', 'iM증권 기업분석(2026.04.15) p.4, 위닝펀드'),
    ('빅뱅 투어 ATP × 모객 매트릭스', 'm40_bigbang_atp', '하나증권(2026.05.11) + 본 리포트, 위닝펀드'),
    ('19개 증권사 컨센 2026E vs 2027E 영업이익 비교', 'consensus_2027', 'FnGuide 19개사 컨센, WiseReport, 위닝펀드'),
    ('K-POP 4사 세대별 IP 보유 매트릭스', 'm45_4comp_gen_ip', '본 리포트 자체 정리, 4사 공시, 위닝펀드'),
    ('보이그룹 파이프라인 6년 공백 시각화', 'm46_boygroup_6y', 'YG 사업보고서, 위닝펀드'),
    # --- 무차입·재평가 (자체) ---
    ('K-POP 4사 순현금/시총 비율 비교', 'cash_ratio', 'FnGuide, KIS API, 위닝펀드'),
    ('4팩터 Z-Score 막대 그래프 (s1_per 1.37', 'z_score_4factor', 'dacon-skills-dashboard 4팩터, 위닝펀드'),
    ('YG 5년 12MF PER 추이 + 역사적 하단 마킹', 'per_historical', 'FnGuide, FDR, 위닝펀드'),
    # --- 리스크 (자체) ---
    ('블랙핑크 그룹 vs 개별 활동 시기 도식', 'm58_blackpink_solo', 'YG 공시·아티스트 개별 소속사, 위닝펀드'),
    ('9월 보이그룹 데뷔 결과별 적정주가 시나리오', 'm62_sept_scenario', '본 리포트 자체 시나리오, 위닝펀드'),
    # --- 재무 (자체) ---
    ('1Q26 부문별 매출 분해 (상제품 652.7', 'q1_revenue', 'WiseReport 1Q26 잠정(2026.05.08), 위닝펀드'),
    ('외국인·기관·개인 20일 순매수 막대', 'supply_20d', 'KIS API investor_trend, 위닝펀드'),
    ('5년 매출·영업이익·OPM 시계열 막대', 'income_5y', 'FnGuide SVD_Finance(연결), 위닝펀드'),
    ('ROE 듀폰 분해 5년 시계열', 'm69_roe_dupont', 'FnGuide SVD_Finance, FnGuide 컨센, 위닝펀드'),
    # --- 밸류에이션 (자체) ---
    ('매출 부문별 5년 시계열 + 2028E 추정', 'm72_revenue_2028e', 'FnGuide + 본 리포트 추정, 위닝펀드'),
    ('YG 5년 발표기준 OPM 시계열 + 2028E 16%', 'm73_opm_2028e', 'WiseReport 5년 OPM + 본 리포트, 위닝펀드'),
    ('2028E EPS 8단계 분해 워터폴', 'eps_waterfall', '본 리포트 자체 추정, FnGuide 컨센, 위닝펀드'),
    ('Target PER 24배 덧셈 워터폴', 'per_waterfall', '본 리포트 자체 산정, FnGuide 4사 컨센, 위닝펀드'),
    ('4사 Implied Target PER 비교 (하이브 22배', 'm76_4comp_implied_per', 'WiseReport 4사 컨센 + 본 리포트, 위닝펀드'),
    ('보수 가정 5가지 + 적정주가 민감도 매트릭스', 'm79_5_assumptions', '본 리포트 자체 민감도, 위닝펀드'),
    ('Bear/Base/Bull 3시나리오 적정주가 매트릭스', 'scenario_3', '본 리포트 자체 시나리오, 위닝펀드'),
    ('순현금 Bridge 4년 시계열 (2025A 2,737', 'cash_bridge', 'FnGuide 컨센 + 본 리포트 추정, 위닝펀드'),
    ('IP × 분기 매출 시간표 매트릭스', 'm83_ip_quarterly', '본 리포트 IP 라이프사이클, 위닝펀드'),
    ("'양성 시스템 검증 매트릭스'", 'm86_system_validation', '본 리포트 자체 정리, 위닝펀드'),
    ('가격 단계적 상승 사다리', 'm88_price_ladder', '본 리포트 자체 산정, 위닝펀드'),
    ('BUY 결론 인포그래픽', 'm89_buy_conclusion', '본 리포트 자체 산정, 위닝펀드'),
    # 제거된 부정확 매핑 (모두 텍스트 메모로 되돌림):
    # - m39_objectdanga (유튜브 도표가 잡힘, 객단가 X)
    # - m42_bigbang_10y (4개 도표 모음, 분리 어려움)
    # - m60_tencent_pl (재무제표 잡힘)
    # - m64_op_vs_ni (키움 P/E Band 잡힘)
    # - m65_foreign_ratio (매출 표 잡힘)
    # - m70_roe_recovery (주가 추이 잡힘)
    # - m71_consensus_dist (재무제표 잡힘)
    # - m77_report_vs_consensus (음반 판매량 잡힘)
    # - m_mirae_per_pbr_band (페이지 일부)
    # - m_mirae_quarterly (검증 필요)
]


def find_match(memo, used_pngs):
    for kw, png, src in EXACT:
        if kw in memo and png not in used_pngs:
            p = CHARTS / f'{png}.png'
            if p.exists():
                used_pngs.add(png)
                return p, kw, src
    return None, None, None


def short(memo):
    t = memo.replace('📊 [도표 제안]', '').strip()
    t = re.sub(r'^(★\s*)?(핵심\s*도표\s*권장|추가\s*권장)\s*:\s*', '', t)
    t = re.sub(r'\s*\[(스크린샷|신규).*?\].*$', '', t)
    t = re.sub(r'\s*—\s*이미 적용됨\s*$', '', t)
    return t.strip()


def main():
    doc = Document(str(DOC))
    memos = [(i, p) for i, p in enumerate(doc.paragraphs) if '📊 [도표 제안]' in p.text]
    print(f'도표 제안: {len(memos)}개\n')
    used = set()
    inserted = 0; text_only = 0
    for idx, p in memos:
        png, kw, src = find_match(p.text, used)
        cap = short(p.text)
        if png:
            for r in list(p.runs): r._r.getparent().remove(r._r)
            r1 = p.add_run(); r1.text = f'[도표] {cap}\n'
            r1.font.size = Pt(8); r1.font.bold = True
            r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x6D); r1.font.italic = True
            r2 = p.add_run(); r2.add_picture(str(png), width=Inches(5.5))
            r3 = p.add_run(); r3.text = f'\n자료: {src}'
            r3.font.size = Pt(7); r3.font.color.rgb = RGBColor(0x88, 0x88, 0x88); r3.font.italic = True
            inserted += 1
        else:
            for r in list(p.runs): r._r.getparent().remove(r._r)
            r = p.add_run(); r.text = f'[도표 권장] {cap}'
            r.font.size = Pt(8); r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA); r.font.italic = True
            text_only += 1
    print(f'정밀 1:1 PNG 삽입: {inserted}개')
    print(f'텍스트 메모 유지: {text_only}개')
    print(f'사용된 고유 PNG: {len(used)}개')
    doc.save(str(DOC))
    print(f'저장: {DOC.stat().st_size/1024:.1f} KB')


if __name__ == '__main__':
    main()
