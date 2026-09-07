"""
v12 — 양식 그래픽 100% 보존 + 본문·도표만 교체 → PPT → PDF (PowerPoint COM)
- 양식: 1슬라이드 = 본문 2 + 도표 4 (좌우 분할, 각 H=3.26)
- 도표가 4분할이라 자동으로 작아짐 (사용자 우려 해결)
- 워드 슬라이드 1:1 매핑 (29장)
"""
import sys, io, json, shutil, tempfile, subprocess
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from copy import deepcopy
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pathlib import Path
from PIL import Image

# 우수 리포트 양식 정확 색상 (클래시스 p5 RGB 분석값)
TPL_BLUE = RGBColor(0x26, 0x83, 0xC6)  # 양식 소제목·헤더 정확한 파란색
NAVY = RGBColor(0x1E, 0x3A, 0x6D)
RED = RGBColor(0xC8, 0x10, 0x2E)
GREY = RGBColor(0x44, 0x44, 0x44)
BLACK = RGBColor(0x00, 0x00, 0x00)

SRC = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_기업분석 리포트 양식_clean.pptx')
DST = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_v12.pptx')
PDF_DST = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\18-1_와이지엔터테인먼트_PDF_v12.pdf')
BOXES_JSON = Path('scripts/_yg_v10_boxes.json')
CHARTS = Path(r'C:\Users\hsh\Desktop\vibecoding\주식 ai 리서치 리포트 에이전트\output\와이지엔터\_charts')


def add_blank_slide_from_template(dst, src_slide):
    layout = src_slide.slide_layout
    new_slide = dst.slides.add_slide(layout)
    for sh in src_slide.shapes:
        new_el = deepcopy(sh.element)
        new_slide.shapes._spTree.insert_element_before(new_el, 'p:extLst')
    return new_slide


def set_text(tf, text, font_size=None, bold=None, color=None):
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    if font_size: r.font.size = Pt(font_size)
    if bold is not None: r.font.bold = bold
    if color: r.font.color.rgb = color


def write_body(tf, head, body, font_size=11):
    """우수 리포트 양식 정확 — 소제목 13pt Bold TPL_BLUE + 본문 11pt 검은색"""
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if head:
        r = p.add_run(); r.text = head
        r.font.size = Pt(13); r.font.bold = True
        r.font.color.rgb = TPL_BLUE  # 양식 정확한 파란색
        r.font.name = '맑은 고딕'
        for line in body.split('\n'):
            if not line.strip(): continue
            p = tf.add_paragraph()
            r = p.add_run(); r.text = line
            r.font.size = Pt(font_size); r.font.color.rgb = BLACK
            r.font.name = '맑은 고딕'
    else:
        r = p.add_run(); r.text = body
        r.font.size = Pt(font_size); r.font.color.rgb = BLACK
        r.font.name = '맑은 고딕'


def find_text_shape(slide, anchor_text, top_min=0, top_max=12):
    """양식의 placeholder 텍스트박스 찾기"""
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        top_in = Emu(sh.top).inches if sh.top else 0
        if anchor_text in t and top_min <= top_in <= top_max:
            return sh
    return None


def find_group_shapes(slide):
    """도표 placeholder GROUP 위치 측정 (4개)"""
    groups = []
    for sh in slide.shapes:
        if sh.shape_type == 6:  # GROUP
            l = Emu(sh.left).inches if sh.left else 0
            t = Emu(sh.top).inches if sh.top else 0
            w = Emu(sh.width).inches if sh.width else 0
            h = Emu(sh.height).inches if sh.height else 0
            groups.append((sh, l, t, w, h))
    groups.sort(key=lambda x: (x[2], x[1]))  # top → left 순
    return groups


def remove_chart_groups(slide):
    """양식 도표 placeholder GROUP만 제거 (LINE/로고 등 보존)"""
    groups = find_group_shapes(slide)
    for sh, _, _, _, _ in groups:
        sh._element.getparent().remove(sh._element)


def find_chart_caption_textbox(slide, position_idx):
    """양식 도표 캡션 텍스트박스(도표 00. 제목) 4개 — position_idx (0~3) 반환"""
    candidates = []
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if '도표 00' in t or t.strip().startswith('도표'):
            l = Emu(sh.left).inches if sh.left else 0
            top = Emu(sh.top).inches if sh.top else 0
            candidates.append((sh, l, top))
    candidates.sort(key=lambda x: (x[2], x[1]))
    if position_idx < len(candidates):
        return candidates[position_idx][0]
    return None


def find_source_textbox(slide, position_idx):
    """양식 도표 출처 텍스트박스(자료 : OO증권) 4개"""
    candidates = []
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if t.startswith('자료 :') or t.startswith('자료:'):
            l = Emu(sh.left).inches if sh.left else 0
            top = Emu(sh.top).inches if sh.top else 0
            candidates.append((sh, l, top))
    candidates.sort(key=lambda x: (x[2], x[1]))
    if position_idx < len(candidates):
        return candidates[position_idx][0]
    return None


def insert_chart_in_group_area(slide, group_info, png_path):
    """GROUP 위치에 PNG 삽입 (비율 fit + 중앙)"""
    _, l, t, w, h = group_info
    if not png_path.exists(): return
    try:
        im = Image.open(str(png_path))
        iw, ih = im.size
        if w * (ih / iw) <= h:
            actual_w = w
            actual_h = w * (ih / iw)
            off_l = l
            off_t = t + (h - actual_h) / 2
        else:
            actual_h = h
            actual_w = h * (iw / ih)
            off_l = l + (w - actual_w) / 2
            off_t = t
        slide.shapes.add_picture(str(png_path), Inches(off_l), Inches(off_t),
                                 Inches(actual_w), Inches(actual_h))
    except Exception as e:
        print(f'  ! 도표 삽입 실패: {png_path}: {e}')


# ============== 워드 슬라이드 → PPT 매핑 ==============
# 각 워드 슬라이드 = (양식 idx, 박스 리스트, 도표 매핑 리스트)

CAT_TEMPLATE = {
    '산업': 3, '기업': 4, '기업-주가': 5,
    '투자': 6, '리스크': 8, '재무': 9, '밸류': 10,
}


def categorize(slide_num):
    n = int(slide_num) if str(slide_num).isdigit() else 0
    if 3 <= n <= 6: return '산업'
    if 7 <= n <= 12: return '기업'
    if n == 13: return '기업-주가'
    if n == 14: return '기업'
    if 15 <= n <= 19: return '투자'
    if 20 <= n <= 21: return '리스크'
    if 22 <= n <= 25: return '재무'
    if 26 <= n <= 28: return '밸류'
    return '기업'


def short_copy(box_title):
    t = box_title.split('—')[0].split(':')[0].strip()
    return t[:24]


# 워드 슬라이드 num → 도표 매핑 (최대 4개)
# (PNG 파일명, 도표 번호, 도표 제목, 출처)
CHART_PER_SLIDE = {
    3: [('industry_consensus_vs_price', '도표 01.', '엔터 4사 영업이익 컨센 +6% vs 주가 -39%', '키움증권, NH투자증권')],
    4: [('industry_ifpi', '도표 02.', 'IFPI 글로벌 음악산업 매출 11년', 'IFPI 2025')],
    5: [('industry_triggers', '도표 03.', '회복 트리거 4가지', '키움증권(2026.05.07)'),
        ('industry_4comp', '도표 04.', 'K-POP 4사 12MF PER + 시총', 'FnGuide, WiseReport')],
    6: [('m45_4comp_gen_ip', '도표 05.', 'K-POP 4사 세대별 IP 분산', 'WiseReport, 분기보고서')],
    7: [('analyst_iM_p2', '도표 06.', 'YG 5년 사업·실적 구조', 'iM증권(2026.04.15)')],
    8: [('company_revenue_5y', '도표 07.', 'YG 매출 부문별 5년 (공연 7배)', 'FnGuide, DART')],
    9: [('m_first5_albums', '도표 08.', 'YG IP 첫 5개 앨범 — 베몬 25개월', 'DART 사업보고서')],
    10: [('baemon_youtube', '도표 09.', '베몬 YT 1,200만 — K팝 걸그룹 최단', 'iM증권(2026.04.15)'),
         ('company_comeback_cycle', '도표 10.', 'YG IP 평균 컴백 간격', 'DART')],
    11: [('m45_4comp_gen_ip', '도표 11.', 'K-POP 4사 두 문제 동시 해결', 'WiseReport')],
    12: [('analyst_교보_p2', '도표 12.', '9월 신인 보이그룹·NEXT MONSTER 일정', '교보증권')],
    13: [('m30_5y_events', '도표 13.', '5년 주가 9개 이벤트 마킹', 'FDR, 공시'),
         ('pattern_190', '도표 14.', '+190% 패턴 (FDR 일봉 실측)', 'FinanceDataReader')],
    14: [('ip_timeline', '도표 15.', '5종 IP 동시 가동 타임라인', 'DART, 공시'),
         ('m35_cycle_compare', '도표 16.', '24.10 vs 26.08 사이클 비교', 'FDR, 본 리포트')],
    15: [('m38_bigbang_tour', '도표 17.', '빅뱅 20주년 14개 도시 투어', 'iM증권'),
         ('consensus_2027', '도표 18.', '컨센 2027E OP 764억 — 빅뱅 미반영', 'FnGuide 19개사')],
    16: [('m46_boygroup_6y', '도표 19.', '6년 만의 보이그룹 — 9월 5인조', 'YG 2026 PLAN')],
    17: [('analyst_SK_p3', '도표 20.', 'YG 6종 IP 분산 효과', 'SK증권(2026.05.11)')],
    18: [('cash_ratio', '도표 21.', 'K-POP 4사 순현금/시총', 'FnGuide')],
    19: [('z_score_4factor', '도표 22.', '4팩터 Z-Score — 303종목 1등', 'dacon-skills-dashboard'),
         ('per_historical', '도표 23.', '12MF PER 15.9 — 역사적 하단', 'FnGuide, FDR')],
    20: [('analyst_한화_p3', '도표 24.', '리스크 1·2·3 정리', '한화투자증권(2026.04.13)')],
    21: [('m62_sept_scenario', '도표 25.', '9월 신인 데뷔 실패 시 시나리오', '본 리포트')],
    22: [('q1_revenue', '도표 26.', '1Q26 부문별 매출 (+46.9%)', 'WiseReport 1Q26')],
    23: [('supply_20d', '도표 27.', '20일 수급 — 바닥 분산', 'KIS API')],
    24: [('income_5y', '도표 28.', 'YG 5년 매출·영업이익', 'FnGuide')],
    25: [('m69_roe_dupont', '도표 29.', 'ROE 듀폰 분해 5년', 'FnGuide')],
    26: [('per_waterfall', '도표 30.', 'Target PER 24배 워터폴', '본 리포트, FnGuide'),
         ('eps_waterfall', '도표 31.', '2028E EPS 8단계 분해', '본 리포트')],
    27: [('m79_5_assumptions', '도표 32.', '5가지 보수 가정 매트릭스', '본 리포트'),
         ('cash_bridge', '도표 33.', '순현금 Bridge 4년', 'FnGuide, 본 리포트'),
         ('m83_ip_quarterly', '도표 34.', 'IP 라이프사이클', '본 리포트')],
    28: [('scenario_3', '도표 35.', 'Bear/Base/Bull 3시나리오', '본 리포트'),
         ('m88_price_ladder', '도표 36.', '가격 사다리', '본 리포트')],
}


def fill_template_slide(slide, header, slide_boxes, charts):
    """양식 그대로 + 본문 + 도표 채우기"""
    # 1. 헤더 (양식 상단 좌측)
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        l = Emu(sh.left).inches if sh.left else 0
        t = Emu(sh.top).inches if sh.top else 0
        txt = sh.text_frame.text.strip()
        if t < 0.5 and l < 1.0 and len(txt) > 3:
            set_text(sh.text_frame, header, font_size=20, bold=True, color=NAVY)

    # 2. 본문 박스 1·2 채우기
    # 양식의 본문 박스 위치: 상단 T=0.83~1.82, 하단 T=5.86~6.85
    body_top = None
    body_bot = None
    copy_top = None
    copy_bot = None
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        l = Emu(sh.left).inches if sh.left else 0
        t = Emu(sh.top).inches if sh.top else 0
        w = Emu(sh.width).inches if sh.width else 0
        txt = sh.text_frame.text

        if l >= 1.5 and 0.5 < t < 2.5 and w > 4.0:
            body_top = sh
        elif l >= 1.5 and 5.0 < t < 7.0 and w > 4.0:
            body_bot = sh
        elif l < 1.5 and 0.8 < t < 1.8 and ('요약 제목' in txt or '소제목' in txt):
            copy_top = sh
        elif l < 1.5 and 5.5 < t < 6.3 and '요약' in txt:
            copy_bot = sh

    # 박스 채우기 — 우수 리포트 양식 폰트/색
    if len(slide_boxes) >= 1 and body_top:
        box = slide_boxes[0]
        body = '\n'.join(s['body'] for s in box['subs'])
        # 본문 길이에 따라 폰트 크기 (양식 기본 11pt)
        fs = 11 if len(body) < 1200 else 10 if len(body) < 2000 else 9.5
        write_body(body_top.text_frame, box['box_title'], body, font_size=fs)
        if copy_top:
            set_text(copy_top.text_frame, short_copy(box['box_title']),
                    font_size=9, bold=True, color=BLACK)

    if len(slide_boxes) >= 2 and body_bot:
        box = slide_boxes[1]
        body = '\n'.join(s['body'] for s in box['subs'])
        fs = 11 if len(body) < 1200 else 10 if len(body) < 2000 else 9.5
        write_body(body_bot.text_frame, box['box_title'], body, font_size=fs)
        if copy_bot:
            set_text(copy_bot.text_frame, short_copy(box['box_title']),
                    font_size=9, bold=True, color=BLACK)
    elif body_bot:
        # 본문 박스 2가 비어있으면 placeholder 제거
        if '소제목을 입력' in body_bot.text_frame.text:
            set_text(body_bot.text_frame, '', font_size=9)
        if copy_bot and '요약' in copy_bot.text_frame.text:
            set_text(copy_bot.text_frame, '', font_size=9)

    # 3. 박스 3+ 본문이 더 있으면 body_bot에 합쳐 추가
    if len(slide_boxes) >= 3 and body_bot:
        extra = '\n\n'.join([f"[{b['box_title']}]\n" + '\n'.join(s['body'] for s in b['subs'])
                            for b in slide_boxes[2:]])
        cur = body_bot.text_frame.text
        # 박스 2 본문 끝에 박스 3 추가
        p = body_bot.text_frame.add_paragraph()
        r = p.add_run(); r.text = '\n' + extra
        r.font.size = Pt(9); r.font.color.rgb = GREY

    # 4. 도표 영역 — GROUP 위치 측정 + 제거 + PNG 삽입
    groups = find_group_shapes(slide)
    print(f'    도표 자리: {len(groups)}개, 매핑 도표: {len(charts)}개')
    for i, chart in enumerate(charts):
        if i >= len(groups): break
        fname, no, title, src = chart
        png = CHARTS / f'{fname}.png'
        if png.exists():
            insert_chart_in_group_area(slide, groups[i], png)
            # 캡션 + 출처 텍스트박스 교체
            caption_tb = find_chart_caption_textbox(slide, i)
            if caption_tb:
                # 양식 도표 캡션 — 맑은 고딕 Bold 9pt 검은색 (양식 정확)
                set_text(caption_tb.text_frame, f'{no} {title}', font_size=9, bold=True, color=BLACK)
            source_tb = find_source_textbox(slide, i)
            if source_tb:
                # 양식 출처 — 7pt 검은색
                set_text(source_tb.text_frame, f'자료 : {src}', font_size=7, color=BLACK)
    # GROUP placeholder 제거 (도표 자리 그래픽 비우기)
    remove_chart_groups(slide)
    # 사용 안 된 캡션·출처는 비움
    for i in range(len(charts), 4):
        caption_tb = find_chart_caption_textbox(slide, i)
        if caption_tb: set_text(caption_tb.text_frame, '', font_size=8)
        source_tb = find_source_textbox(slide, i)
        if source_tb: set_text(source_tb.text_frame, '', font_size=8)


def update_page_no(prs):
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if not sh.has_text_frame: continue
            t = sh.text_frame.text.strip()
            l = Emu(sh.left).inches if sh.left else 0
            top = Emu(sh.top).inches if sh.top else 0
            if l > 7 and top > 11 and ('|' in t or t.strip().isdigit()):
                set_text(sh.text_frame, f'|   {i+1}', font_size=9, color=GREY)


def set_cell(cell, text, font_size=None, bold=None, color=None):
    tf = cell.text_frame
    if not tf.paragraphs: return
    p = tf.paragraphs[0]
    if p.runs:
        first = p.runs[0]
        first.text = text
        if font_size: first.font.size = Pt(font_size)
        if bold is not None: first.font.bold = bold
        if color: first.font.color.rgb = color
        for r in list(p.runs[1:]): r._r.getparent().remove(r._r)
    else:
        r = p.add_run(); r.text = text
        if font_size: r.font.size = Pt(font_size)
        if bold is not None: r.font.bold = bold
        if color: r.font.color.rgb = color
    for pp in list(tf.paragraphs[1:]): pp._p.getparent().remove(pp._p)


def fill_cover(slide, cover_boxes):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if '회사명' in t and '종목코드' in t:
            set_text(sh.text_frame, '와이지엔터테인먼트 (122870)', font_size=28, bold=True, color=NAVY)
        elif t == '리포트 제목':
            set_text(sh.text_frame, '5년 +190% 패턴, 다시 그릴 차례', font_size=16, color=GREY)
        elif '18-1 Equity Research' in t:
            set_text(sh.text_frame, '18-1 Equity Research', font_size=9)
        elif '2026.05.00' in t or '2026.00.00' in t:
            set_text(sh.text_frame, '2026.05.27.', font_size=9)
        elif 'Winning Fund 17-1' in t:
            set_text(sh.text_frame, 'Winning Fund 18-1 Research Report', font_size=9)
        elif '최종 제출 시' in t or '안내사항' in t:
            tf = sh.text_frame; tf.clear()
            first = True
            for box in cover_boxes:
                title = box['box_title'].split(' ', 1)[0]
                if first: p = tf.paragraphs[0]; first = False
                else: p = tf.add_paragraph()
                r = p.add_run(); r.text = title
                r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = NAVY
                body = '\n'.join(s['body'] for s in box['subs'])
                for line in body.split('\n'):
                    if not line.strip(): continue
                    p = tf.add_paragraph()
                    r = p.add_run(); r.text = line
                    r.font.size = Pt(8); r.font.color.rgb = GREY

    # 테이블
    for sh in slide.shapes:
        if sh.shape_type != 19: continue
        t_pos = Emu(sh.top).inches if sh.top else 0
        tbl = sh.table
        rows, cols = len(tbl.rows), len(tbl.columns)
        if 5.5 < t_pos < 6.0 and rows == 8 and cols == 3:
            data = [
                ('Stock Data', '26.05.27.', ''),
                ('KOSDAQ지수(pt)', '', '789'),
                ('52주 최고/최저(원)', '', '79,000/47,000'),
                ('거래량(천주)/거래대금(억)', '', '215/103'),
                ('시가총액(억원)', '', '8,981'),
                ('발행주식수(천 주)', '', '18,691'),
                ('주요주주 지분율', '', '22.95%'),
                ('외국인 지분율', '', '10.86%'),
            ]
            for ri, row in enumerate(data):
                for ci, val in enumerate(row):
                    set_cell(tbl.cell(ri, ci), val, font_size=8)
        elif 7.0 < t_pos < 7.5:
            try:
                set_cell(tbl.cell(0, 0), 'Consensus Data', font_size=8, bold=True)
                if cols >= 3:
                    set_cell(tbl.cell(0, 1), '2025', font_size=8, bold=True)
                    set_cell(tbl.cell(0, 2), '2026E', font_size=8, bold=True)
                cd = [('매출액(억원)', '5,454', '6,010'), ('영업이익(억원)', '522', '806'),
                      ('순이익(억원)', '537', '684'), ('EPS(원)', '1,974', '3,108')]
                for ri, row in enumerate(cd):
                    if ri+1 >= rows: break
                    for ci, val in enumerate(row):
                        if ci >= cols: break
                        set_cell(tbl.cell(ri+1, ci), val, font_size=8)
            except: pass
        elif 8.3 < t_pos < 8.7:
            try:
                set_cell(tbl.cell(0, 0), '수익률', font_size=8, bold=True)
                if cols >= 4:
                    for ci, h in enumerate(['1M', '6M', '12M']):
                        set_cell(tbl.cell(0, ci+1), h, font_size=8, bold=True)
                    if rows >= 2:
                        set_cell(tbl.cell(1, 0), '절대주가(%)', font_size=8)
                        for ci, v in enumerate(['-3.8', '-26.82', '-24.22']):
                            set_cell(tbl.cell(1, ci+1), v, font_size=8)
            except: pass
        elif 2.8 < t_pos < 3.1:
            try:
                set_cell(tbl.cell(0, 0), '종목정보', font_size=10, bold=True, color=NAVY)
                info = [('업종', 'KOSDAQ / 엔터'), ('주가 (5/27)', '48,050원'), ('상장일', '2011.11.23')]
                for ri, row in enumerate(info):
                    if ri+1 >= rows: break
                    for ci, val in enumerate(row):
                        if ci >= cols: break
                        set_cell(tbl.cell(ri+1, ci), val, font_size=8)
            except: pass
        elif 3.8 < t_pos < 4.2:
            try:
                if rows >= 1: set_cell(tbl.cell(0, 0), 'BUY', font_size=28, bold=True, color=NAVY)
                pairs = [('목표 주가(원)', '89,000', NAVY), ('현재 주가(원)', '48,050', None),
                         ('상승 여력(%)', '+85.2', RED)]
                for ri, (lbl, val, c) in enumerate(pairs):
                    if ri+1 >= rows: break
                    set_cell(tbl.cell(ri+1, 0), lbl, font_size=9)
                    if cols >= 2:
                        set_cell(tbl.cell(ri+1, 1), val, font_size=12, bold=True, color=c or NAVY)
            except: pass


def main():
    print('=== v12 빌드 (양식 그래픽 100% 보존) ===\n')

    with open(BOXES_JSON, encoding='utf-8') as f:
        boxes = json.load(f)
    cover_boxes = [b for b in boxes if b['slide_num'] == '1']

    # 워드 슬라이드 num별로 박스 그룹화
    slide_groups = {}
    for b in boxes:
        if b['slide_num'] in ('1', '2'): continue
        sn = b['slide_num']
        slide_groups.setdefault(sn, []).append(b)

    # 워드 슬라이드 순서 (3~28)
    slide_nums = sorted(slide_groups.keys(), key=lambda x: int(x))
    print(f'워드 슬라이드: {len(slide_nums)}장 (Cover/Index/재무제표 제외)')
    print(f'박스 분포: {[(sn, len(slide_groups[sn])) for sn in slide_nums]}')

    src = Presentation(str(SRC))

    # dst 초기화
    tmp = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False); tmp.close()
    shutil.copy(SRC, tmp.name)
    dst = Presentation(tmp.name)
    sldIdLst = dst.slides._sldIdLst
    while len(list(sldIdLst)) > 0:
        sld = list(sldIdLst)[0]
        rId = sld.get(qn('r:id'))
        dst.part.drop_rel(rId)
        sldIdLst.remove(sld)

    # 1. Cover
    add_blank_slide_from_template(dst, src.slides[0])
    # 2. Index
    add_blank_slide_from_template(dst, src.slides[1])
    # 3~. 워드 슬라이드별 본문
    for sn in slide_nums:
        cat = categorize(sn)
        tmpl_idx = CAT_TEMPLATE[cat]
        add_blank_slide_from_template(dst, src.slides[tmpl_idx])
    # 마지막: 재무제표
    add_blank_slide_from_template(dst, src.slides[11])

    total = len(dst.slides)
    print(f'\n빌드: {total}장\n')

    CATEGORY_HEADER = {
        '산업': '산업분석', '기업': '기업분석', '기업-주가': '기업분석',
        '투자': '투자포인트', '리스크': '리스크', '재무': '재무분석', '밸류': '밸류에이션',
    }

    # Cover
    print('Cover 채우기...')
    fill_cover(dst.slides[0], cover_boxes)

    # 본문 (Slide 3~)
    for i, sn in enumerate(slide_nums):
        slide_idx = i + 2  # Cover=0, Index=1
        slide = dst.slides[slide_idx]
        cat = categorize(sn)
        header = CATEGORY_HEADER[cat]
        slide_boxes = slide_groups[sn]
        charts = CHART_PER_SLIDE.get(int(sn), [])
        print(f'Slide {slide_idx+1:2}: {header} | 박스 {len(slide_boxes)}개 | 도표 {len(charts)}개  ({sn})')
        fill_template_slide(slide, header, slide_boxes, charts)

    update_page_no(dst)

    dst.save(str(DST))
    print(f'\nPPT 저장: {DST.name} ({DST.stat().st_size / 1024:.1f} KB / {total}장)')

    # PPT → PDF 변환 (PowerShell COM, Visible 설정 없이)
    print('\n=== PPT → PDF 변환 (PowerPoint COM) ===')
    PDF_DST.unlink(missing_ok=True)
    ps_script = f'''try {{
    $ppt = New-Object -ComObject PowerPoint.Application
    Start-Sleep -Milliseconds 500
    $pres = $ppt.Presentations.Open("{DST}", $true, $false, $false)
    Start-Sleep -Milliseconds 500
    $pres.SaveAs("{PDF_DST}", 32)
    $pres.Close()
    $ppt.Quit()
    Write-Output "OK"
}} catch {{
    Write-Output "ERR: $_"
}}
'''
    tmp_ps = tempfile.NamedTemporaryFile(suffix='.ps1', delete=False, mode='w', encoding='utf-8-sig')
    tmp_ps.write(ps_script)
    tmp_ps.close()
    result = subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', tmp_ps.name],
                           capture_output=True, text=True, encoding='cp949', timeout=120)
    Path(tmp_ps.name).unlink(missing_ok=True)
    print(f'PS stdout: {result.stdout[:300]}')
    if PDF_DST.exists():
        print(f'PDF 저장: {PDF_DST.name} ({PDF_DST.stat().st_size / 1024:.1f} KB)')
    else:
        print(f'PDF 변환 실패. stderr:\n{result.stderr[:500]}')


if __name__ == '__main__':
    main()
